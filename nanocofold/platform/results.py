"""Read-only adapters for legacy and current runs. No analysis is rerun here."""
from __future__ import annotations

import csv
import copy
import hashlib
import io
import json
import math
import re
import zipfile
import os
import tempfile
from pathlib import Path

import yaml

from ..config import ASSET_ROOT, resolve_run_artifact
from ..binding_site_assessment import hotspot_filter_fraction
from ..utils.metric_provenance import metric_provenance
from .errors import PublicIssues, public_error_data
from .paths import contained, native, sha256

METRICS = ("ipsae", "iptm", "interface_plddt", "hotspot_contact_frac", "n_clashes", "interface_pae", "dev_severe")
GEOMETRY_METRICS = ("n_hbonds", "n_salt_bridges", "n_clashes")
GEOMETRY_METADATA = tuple(metric + suffix for metric in GEOMETRY_METRICS
                          for suffix in ("_method", "_unit", "_cutoff_angstrom"))
ADDITIONAL_METRICS = ("ptm", "pdockq", "pdockq_method", "pdockq2", "pdockq2_target", "pdockq2_binder",
                      "pdockq_status", "pdockq_n_contacts", "pdockq_raw_author_value",
                      "pdockq2_n_contacts", "pdockq2_method", "pdockq2_status",
                      "n_hbonds", "n_salt_bridges", *GEOMETRY_METADATA,
                      "dev_flags", "esm_score", "antifold",
                      "interface_pae_ab", "interface_pae_ba", "interface_pae_median",
                      "interface_pae_p90", "target_plddt", "binder_plddt", "complex_plddt")
SEARCH_FIELDS = ("candidate_id", "sequence", "exclusion_reason", "exclusion_reasons", "error")
TABLES = ("03_score/scored_candidates.csv", "02_screen/candidates.csv", "02_design/design_candidates.csv",
          "02_mature/esm2_scored_library.csv", "02_mature/prefiltered_candidates.csv", "02_design/all_generated_candidates.csv")
METRIC_HELP = {
    "ipsae": "PAE 기반 구조 일관성 점수입니다. VHH 모드는 선택한 표적–VHH 계면이며, 일반 복합체 요약은 ipsae_scope·ipsae_chain_a·ipsae_chain_b에 기록된 체인 쌍을 확인하세요. 결합 확률·KD가 아닙니다.",
    "iptm": "VHH 모드에서는 선택한 표적–VHH 쌍 신뢰도입니다. 일반 복합체 모델 요약의 전역 ipTM과 구분하고 iptm_source·iptm_scope를 확인하세요. 실제 결합·KD가 아닙니다.",
    "complex_iptm": "native 전체 복합체 ipTM입니다. 3체인 이상에서 표적–VHH 쌍 ipTM을 대체하거나 쌍 게이트를 통과시키지 않습니다.",
    "iptm_source": "기록된 쌍 ipTM의 native 출처·체인 축 대응·방향별 요약법입니다. unavailable은 미평가를 뜻합니다.",
    "ipsae_scope": "ipSAE 요약에 실제 사용된 체인 쌍·집계 범위입니다. 기록되지 않은 다중 체인 요약은 전체 복합체 점수로 해석하지 않습니다.",
    "ipsae_chain_a": "요약 ipSAE의 첫 체인입니다. 다른 체인 쌍은 pair_metrics에서 확인합니다.",
    "ipsae_chain_b": "요약 ipSAE의 둘째 체인입니다. 다른 체인 쌍은 pair_metrics에서 확인합니다.",
    "iptm_unavailable_reason": "표적–VHH 쌍 ipTM 미평가의 기록 사유입니다. 미평가를 0 또는 통과로 바꾸지 않습니다.",
    "interface_plddt": "계면 잔기의 예측 구조 신뢰도입니다. 실험 구조의 B-factor와 구분하세요.",
    "hotspot_contact_frac": "요청한 결합 부위 중 접촉 거리 조건을 만족한 비율입니다. 구조 대응이 불완전하면 접촉률을 순위에 사용하지 않고 보류합니다.",
    "n_clashes": "고정 거리 기준 이내의 무거운 원자쌍 수입니다. 표적·보존 파트너와 VHH 사이를 합산합니다. 잔기 수나 원소별 반데르발스 충돌 판정이 아닙니다.",
    "n_hbonds": "거리 기준의 N/O 근접 원자쌍 수입니다. 공여·수용체와 각도 검증을 하지 않으므로 확인된 수소결합 수가 아닙니다.",
    "n_salt_bridges": "Lys/Arg 양전하 곁사슬과 Asp/Glu 음전하 곁사슬의 거리 기준 원자쌍 수입니다. His의 양성자화나 중성 Asn/Gln은 가정하지 않습니다. 기존 계산법 미기록 결과는 별도로 읽으세요.",
    "pdockq": "CB 8 Å 접촉(Gly는 CA)과 신뢰도를 사용하는 탐색 지표입니다. 계면 없음·신뢰도 누락은 미평가 상태이며 원저자 반환값은 별도로 기록됩니다.",
    "interface_pae": "계면 접촉 잔기 사이의 예측 정렬 오차입니다. 참조 실험 구조와의 실측 거리 오차가 아닙니다.",
    "pdockq2": "PAE와 계면 pLDDT를 이용한 pDockQ2입니다. 두 방향 중 최소값이며 CA 8 Å 접촉을 사용합니다. AlphaFold-Multimer에서 보정된 지표로 다른 엔진의 값은 탐색용입니다. 실측 DockQ나 KD가 아닙니다.",
    "interface_pae_ab": "접촉 쌍의 PAE 행 표적·열 VHH 평균(Å)입니다. 반대 방향과 다를 수 있습니다.",
    "interface_pae_ba": "접촉 쌍의 PAE 행 VHH·열 표적 평균(Å)입니다.",
    "interface_pae_median": "양방향 접촉 쌍 PAE의 중앙값(Å)입니다.",
    "interface_pae_p90": "양방향 접촉 쌍 PAE의 90백분위(Å)입니다. 높은 오차 구간을 요약합니다.",
    "dev_severe": "free_cys·glycosylation·deamidation_cdr·isomerization_cdr·hydrophobic_patch_cdr의 5종 심한 위험 수입니다. 선별 게이트에 사용하며 전체 위험 dev_flags와 다릅니다.",
    "dev_flags": "평가된 True 위험 플래그의 전체 수입니다. 심한 5종에 oxidation_cdr·pi_out_of_range·charge_out_of_range가 추가되며 종합점수 가중치에 사용됩니다. CDR 미평가는 위험 없음이 아닙니다.",
    "composite_score": "기록된 엔진과 후보 집합 안의 z-score 가중 상대 점수입니다. 적용·제외 가중치는 ranking_policy의 엔진별 정책을 확인하세요. 부분 재평가는 같은 후보의 엔진 성능 비교가 아닙니다. 과거 혼합 순위는 원래 범위를 유지합니다. 선별 순위에는 게이트와 다양성 제한이 추가됩니다. 결합 확률·KD나 검증된 친화도 순위가 아닙니다.",
    "esm_score": "서열/폴딩 적합도 prior입니다. esm_llr가 기록된 성숙화는 단일 치환 ΔLLR를 사용합니다. ΔLLR가 없는 랜덤 성숙화와 design/screen은 예측 후 계산한 평균 로그확률의 배치 내 z-score를 사용합니다. 미평가는 0점이 아니며 실측 결합·친화도 판별값이 아닙니다.",
    "site_selection_eligible": "실행 당시 기록 적격에 지정 결합 부위 접촉을 별도로 확인한 값입니다. 기존 final_eligible와 선별 목록은 보존합니다.",
    "pose_supporting_engines": "표적만 정렬한 뒤 같은 정확한 후보 서열의 자세가 일치하는 서로 다른 엔진 수입니다. 독립 실험 횟수나 점수 평균이 아닙니다.",
    "pose_quality_supporting_engines": "자세가 일치하고 기록 적격·지정 결합 부위를 함께 충족한 서로 다른 엔진 수입니다.",
    "design_binder_rmsd_A": "동일 후보의 생성 구조와 독립 재예측 구조를 표적만 정렬한 후 계산한 VHH Cα RMSD입니다. 기존 VHH 참조 구조와의 RMSD가 아닙니다.",
    "review_rank": "동일 표적·조건·실행 범위 안의 탐색적 후보 검토 순위입니다. 동률을 허용하며 1위도 보류일 수 있습니다. 결합 확률이나 엔진 점수 평균이 아닙니다.",
    "review_priority": "검토 정책 v3는 엔진별 단독 유망 예측과 같은 자세의 다중 엔진 지지를 구분합니다. 두 경우 모두 품질·지정 결합 부위를 확인합니다. 엔진 간 자세 차이와 설계 재현은 별도 진단이며 자동 탈락 조건에 포함되지 않습니다. 결합과 KD는 실험으로 확인하세요.",
}
POSE_ANNOTATIONS = ("site_satisfied", "site_selection_eligible", "pose_cluster_id", "pose_supporting_engines",
                    "pose_quality_supporting_engines", "pose_consensus_status", "design_consistency_status",
                    "design_binder_rmsd_A", "review_priority", "review_rank", "review_rank_scope",
                    "review_priority_reason", "candidate_review_priority", "pose_assessment")


def _geometry_metadata(row):
    """Keep method receipts attached only to their own recorded stage values."""
    for prefix in ("", "colabfold_", "boltz_"):
        for metric in GEOMETRY_METRICS:
            value = row.get(prefix + metric)
            available = (not isinstance(value, bool) and isinstance(value, (int, float))
                         and math.isfinite(value))
            if not available:
                for suffix in ("_method", "_unit", "_cutoff_angstrom"):
                    key = prefix + metric + suffix
                    if key in row:
                        row[key] = None
    return row


def _cell(value, key):
    # Identifiers are opaque strings. A valid FASTA ID such as "false" or
    # "nan" must not become a boolean, missing value or fallback row ID.
    if key.endswith("_id") or key in {"label", "sequence", "target_chain", "binder_chain"}:
        return value if value is not None and value != "" else None
    if value is None or value.strip().lower() in {"", "nan", "null", "none", "inf", "-inf"}:
        return None
    if value.lower() in {"true", "false"}:
        return value.lower() == "true"
    if re.fullmatch(r"[-+]?\d+", value):
        return int(value)
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except ValueError:
        return value


def _exploration(filters):
    """Preserve display state as typed data; never forward UI-only keys to queries."""
    if not isinstance(filters, dict):
        raise ValueError("탐색 조건은 JSON 객체여야 합니다.")
    allowed = {}
    text_limits = {"search": 10000, "sort": 100, "selected_candidate_id": 1000,
                   "selected_observation_id": 2048, "plot_x": 100, "plot_y": 100,
                   "engine_filter": 100}
    for key, limit in text_limits.items():
        if key not in filters:
            continue
        value = filters[key]
        if value is None and key.startswith("selected_"):
            allowed[key] = None
        elif not isinstance(value, str) or len(value) > limit:
            raise ValueError(f"탐색 조건 {key} 형식이 잘못되었습니다.")
        else:
            allowed[key] = value
    for key in ("descending", "eligible"):
        if key in filters:
            if not isinstance(filters[key], bool):
                raise ValueError(f"탐색 조건 {key}는 true/false여야 합니다.")
            allowed[key] = filters[key]
    if "eligibility" in filters:
        value = filters["eligibility"]
        if value not in ("all", "true", "false", "missing"):
            raise ValueError("적격 탐색 조건은 all/true/false/missing 중 하나여야 합니다.")
        allowed["eligibility"] = value
    if "eligibility_basis" in filters:
        if filters["eligibility_basis"] not in ("recorded", "quality"):
            raise ValueError("적격 판정 기준은 recorded/quality 중 하나여야 합니다.")
        allowed["eligibility_basis"] = filters["eligibility_basis"]
    if "hotspot_min" in filters:
        allowed["hotspot_min"] = _hotspot_minimum(filters["hotspot_min"])
    if "result_layers" in filters:
        layers = filters["result_layers"]
        keys = {"target", "binder", "context", "hotspots", "contacts", "clashes", "cdr", "reference"}
        if not isinstance(layers, dict) or not set(layers) <= keys or any(type(v) is not bool for v in layers.values()):
            raise ValueError("구조 레이어 탐색 조건 형식 오류")
        allowed["result_layers"] = layers
    if "viewer_display" in filters:
        display = filters["viewer_display"]
        if not isinstance(display, dict) or not set(display) <= {"sequence", "residues"} or any(type(v) is not bool for v in display.values()):
            raise ValueError("구조 표시 탐색 조건 형식 오류")
        allowed["viewer_display"] = display
    if "selected_residues" in filters:
        labels = filters["selected_residues"]
        if not isinstance(labels, list) or len(labels) > 10000 or any(not isinstance(v, str) or len(v) > 128 for v in labels):
            raise ValueError("선택 잔기 탐색 조건 형식 오류")
        allowed["selected_residues"] = labels
    if "distance_selection" in filters:
        pair = filters["distance_selection"]
        if not isinstance(pair, dict) or set(pair) != {"a", "b"} or any(not isinstance(v, str) or len(v) > 128 for v in pair.values()):
            raise ValueError("거리 선택 탐색 조건 형식 오류")
        allowed["distance_selection"] = pair
    return allowed


def _hotspot_minimum(value):
    if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float))
                              or not math.isfinite(value) or not 0 <= value <= 1):
        raise ValueError("hotspot 최소 접촉률은 0–1 사이의 유한한 숫자여야 합니다.")
    return value


def _assessed(row):
    """Count recorded successful model evaluations, never a contradictory failure."""
    return (row.get("model_record", True) is True and row.get("evaluated") is True
            and not row.get("error") and not row.get("prediction_error")
            and row.get("prediction_status") != "failed")


def _pack_offline_pae(structure, arrays):
    """Losslessly pool PAE bytes; keep original matrix axes on each structure."""
    pae = structure.get("pae")
    if not pae or "values" not in pae:
        return structure
    import base64
    import gzip
    import numpy as np
    matrix = np.asarray(pae["values"], dtype="<f8")
    if (matrix.ndim != 2 or matrix.shape[0] != matrix.shape[1] or not matrix.shape[0]
            or len(pae.get("labels", [])) != matrix.shape[0] or not np.isfinite(matrix).all()
            or (matrix < 0).any()):
        raise ValueError("PAE는 잔기 축과 일치하는 유한한 정사각 행렬이어야 합니다.")
    binary = matrix.tobytes(order="C")
    digest = hashlib.sha256(binary).hexdigest()
    if digest not in arrays:
        arrays[digest] = {"encoding": "gzip-base64-float64le", "dimension": matrix.shape[0],
                          "bytes": len(binary), "sha256": digest,
                          "data": base64.b64encode(gzip.compress(binary, compresslevel=9, mtime=0)).decode("ascii")}
    structure["pae"] = {key: value for key, value in pae.items() if key != "values"}
    structure["pae"]["values_ref"] = digest
    return structure


class ResultReader:
    def __init__(self, run_dir: Path):
        self.root = Path(run_dir).resolve()
        if not self.root.is_dir():
            raise FileNotFoundError("결과 디렉터리가 없습니다.")
        self.issues: list[str] = PublicIssues()
        self.config = self._yaml("run_config.yaml")
        self.state = self._json("run_state.json")
        self.platform_status = self._json("platform_status.json")
        # A verified compact record is authoritative after storage cleanup.
        # Reading it never regenerates scientific values or launches a model.
        self.retained = None
        if (self.root / "retained").exists():
            from .retention import retained_record
            self.retained = public_error_data(retained_record(self.root))
            if self.retained:
                self.config = copy.deepcopy(self.retained["config"])
                self.state = copy.deepcopy(self.retained["state"])
                self.platform_status = copy.deepcopy(self.retained.get("platform_status", {}))
        self.assessment = self._specificity_assessment()

    def _specificity_assessment(self):
        """Verify current derived criteria separately from original prediction settings."""
        if not self.config.get("specificity", {}).get("enabled"):
            return None
        relative = "06_specificity/assessment.json"
        def unknown(reason):
            return {"status": "not_established", "reason": reason, "source": relative,
                    "assessment_config": None, "prediction_config_source": "run_config.yaml"}
        path = self._path(relative)
        if not path.is_file():
            return unknown("assessment_receipt_missing")
        try:
            receipt = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(receipt, dict) or receipt.get("schema") != "specificity_assessment_v1":
                return unknown("assessment_receipt_invalid")
            content = {key: value for key, value in receipt.items() if key != "fingerprint"}
            digest = hashlib.sha256(json.dumps(content, sort_keys=True, default=str).encode()).hexdigest()
            if receipt.get("fingerprint") != digest:
                return unknown("assessment_receipt_invalid")
            prediction = receipt.get("prediction", {})
            config_path = self._path("run_config.yaml")
            if (not isinstance(prediction, dict) or not config_path.is_file()
                    or prediction.get("config_sha256") != sha256(config_path)):
                return unknown("assessment_receipt_source_mismatch")
            outputs = receipt.get("outputs")
            expected = {"03_score/scored_candidates.csv", "06_specificity/comparison.csv",
                        "05_final/selected_candidates.csv", "03_score/ranking_policy.json",
                        "06_specificity/results.json", "06_specificity/observations.csv"}
            if not isinstance(outputs, dict) or set(outputs) != expected:
                return unknown("assessment_receipt_table_mismatch")
            if any(not self._path(name).is_file() or sha256(self._path(name)) != value
                   for name, value in outputs.items()):
                return unknown("assessment_receipt_table_mismatch")
            raw_results = json.loads(self._path("06_specificity/results.json").read_text(encoding="utf-8"))
            if prediction.get("comparison_provenance_fingerprint") != raw_results.get("comparison_provenance_fingerprint"):
                return unknown("assessment_receipt_source_mismatch")
            settings = receipt.get("assessment_config", {})
            policy = json.loads(self._path("03_score/ranking_policy.json").read_text(encoding="utf-8"))
            if (policy.get("assessment_config") != settings or policy.get("method") != receipt.get("method")
                    or receipt.get("method") != "specificity_eligible_lexicographic_then_CDR3_diversity_v1"):
                return unknown("assessment_receipt_settings_invalid")
            thresholds = settings.get("score", {}).get("thresholds")
            maximum = settings.get("rank", {}).get("max_dev_severe")
            final_n = settings.get("specificity", {}).get("final_n")
            finite = lambda value: isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
            if (not isinstance(thresholds, dict) or not all(finite(value) for value in thresholds.values())
                    or not finite(maximum) or type(final_n) is not int or final_n < 1
                    or not re.fullmatch(r"[0-9a-f]{64}", str(receipt.get("source_sha256", "")))):
                return unknown("assessment_receipt_settings_invalid")
            return {"status": "verified", "reason": None, "source": relative,
                    "fingerprint": digest, "assessment_config": settings,
                    "prediction_config_source": "run_config.yaml", "prediction": prediction,
                    "reporting_version": receipt.get("reporting_version"), "source_sha256": receipt["source_sha256"]}
        except (OSError, ValueError, TypeError, AttributeError):
            return unknown("assessment_receipt_invalid")

    def _path(self, relative):
        return contained(self.root, relative)

    def _json(self, relative):
        path = self._path(relative)
        if not path.exists():
            return {}
        try:
            data = json.loads(path.read_text())
            if not isinstance(data, dict):
                raise ValueError("JSON object expected")
            if relative == "run_state.json" and (not isinstance(data.get("phases", {}), dict) or not isinstance(data.get("warnings", []), list)
                                                    or not isinstance(data.get("execution", {}), dict)):
                raise ValueError("상태 필드 형식 오류")
            if relative == "run_state.json":
                if (any(not isinstance(value, dict) for value in data.get("phases", {}).values())
                        or ("state" in data.get("execution", {}) and not isinstance(data["execution"]["state"], str))):
                    raise ValueError("단계·실행 상태 필드 형식 오류")
            if relative == "platform_status.json" and (not isinstance(data.get("state"), str)
                    or data["state"] not in {"cancelled", "interrupted", "failed", "partial_failure", "budget_exhausted", "complete"}
                    or isinstance(data.get("updated"), bool) or not isinstance(data.get("updated"), (int, float)) or not math.isfinite(data["updated"])):
                raise ValueError("플랫폼 상태 필드 형식 오류")
            return public_error_data(data)
        except (ValueError, OSError) as exc:
            self.issues.append(f"{relative}: 읽을 수 없는 JSON ({type(exc).__name__})")
            return {}

    def _yaml(self, relative):
        path = self._path(relative)
        if not path.exists():
            return {}
        try:
            data = yaml.safe_load(path.read_text())
            if not isinstance(data, dict):
                raise ValueError("YAML object expected")
            if any(key in data and not isinstance(data[key], dict) for key in ("run", "target", "screen", "structure", "mature", "score", "rank", "specificity", "complex", "input_review")):
                raise ValueError("설정 필드 형식 오류")
            for section, fields in {"screen": ("comparison", "cascade", "restraint", "colabfold"),
                                    "structure": ("msa",), "score": ("boltz2",),
                                    "mature": ("cdr",)}.items():
                if any(key in data.get(section, {}) and not isinstance(data[section][key], dict) for key in fields):
                    raise ValueError("세부 설정 필드 형식 오류")
            return data
        except (ValueError, yaml.YAMLError, OSError) as exc:
            self.issues.append(f"{relative}: 읽을 수 없는 설정 ({type(exc).__name__})")
            return {}

    def _csv(self, relative):
        path = self._path(relative)
        if not path.exists():
            return []
        try:
            with path.open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, strict=True)
                columns = reader.fieldnames
                if not columns or any(not key for key in columns) or len(set(columns)) != len(columns):
                    raise ValueError("중복·누락 CSV 헤더")
                rows = []
                for row in reader:
                    if None in row or any(value is None for value in row.values()):
                        raise ValueError("CSV 열 개수 불일치")
                    rows.append(public_error_data({key: _cell(value, key) for key, value in row.items()}))
            return rows
        except (ValueError, csv.Error, UnicodeError, OSError) as exc:
            self.issues.append(f"{relative}: 읽을 수 없는 CSV ({exc})")
            return []

    @staticmethod
    def _identifier(row, index):
        return str(row.get("candidate_id") or row.get("design_id") or row.get("variant_id") or row.get("label") or f"row-{index + 1}")

    def all_rows(self):
        from ..result_schema import model_observation_record
        mode = self.config.get("run", {}).get("mode", "unknown")
        if self.assessment:
            self.assessment = self._specificity_assessment()
        if self.retained:
            rows = copy.deepcopy(self.retained["candidates"]["rows"])
            if self.assessment:
                for row in rows:
                    row.update(self._assessment_details(row))
            return [model_observation_record(_geometry_metadata(row), mode=mode) for row in rows]
        if self.config.get("run", {}).get("mode") == "complex":
            recorded = self._json("complex_results.json")
            rows = recorded.get("rows", [])
            if not isinstance(rows, list) or any(not isinstance(r, dict) or not r.get("candidate_id") for r in rows):
                self.issues.append("복합체 모델 목록이 손상되었습니다.")
                return []
            settings = recorded.get("settings", self.config.get("complex", {}))
            if not isinstance(settings, dict):
                settings = {}
            for row in rows:
                row["model_record"] = True
                # General-complex receipts record these conditions at run level.
                # Retain explicit per-model values and keep absent receipts null.
                row.setdefault("msa", settings.get("msa"))
                row.setdefault("restraint", settings.get("restraints"))
                pairs = row.get("pair_metrics")
                if isinstance(pairs, list) and len(pairs) == 1 and isinstance(pairs[0], dict):
                    # A two-chain complex has one unambiguous interface. Show
                    # its recorded metrics without combining multiple pairs.
                    for metric in ("interface_plddt", "interface_pae"):
                        if metric not in row and metric in pairs[0]:
                            row[metric] = pairs[0][metric]
                            row.setdefault("metric_scopes", {})[metric] = {
                                "source": "recorded_pair_metrics",
                                "chains": [pairs[0].get("chain_a"), pairs[0].get("chain_b")]}
                _geometry_metadata(row)
                row["metric_provenance"] = metric_provenance(row)
                row.setdefault("structure_available", False)
                if row["structure_available"]:
                    try:
                        if not row.get("structure_file") or not self._path(row["structure_file"]).is_file():
                            raise FileNotFoundError
                    except (ValueError, OSError):
                        row["structure_available"] = False
                        issue = f"{row['candidate_id']}: 저장된 구조가 누락되었습니다."
                        if issue not in self.issues:self.issues.append(issue)
            return [model_observation_record(row, mode=mode) for row in rows]
        # Begin with all generated sequences; add evaluated metrics without dropping unscored rows.
        records, order = {}, []
        sources = [path for path in TABLES[1:] if self._path(path).exists()]
        comparison = self.config.get("screen", {}).get("comparison", {}).get("enabled") is True
        # Comparison IDs represent independent observations of the same source
        # sequence. Source-stage parent CSVs are provenance, not extra models.
        if comparison:
            sources = []
        if not sources and not self._path(TABLES[0]).exists():
            self.issues.append("후보 CSV를 찾지 못했습니다. 분석이 완료됐는지 확인하세요.")
        for relative in sources + [TABLES[0]]:
            for index, row in enumerate(self._csv(relative)):
                identifier = self._identifier(row, index)
                if identifier not in records:
                    records[identifier] = {}
                    order.append(identifier)
                records[identifier].update(row)
                if relative == TABLES[0]:
                    records[identifier]["model_record"] = True
                    records[identifier]["evaluated"] = row.get("final_evaluated", True) is True
                    if row.get("cif_path") and not row.get("structure_stage"):
                        records[identifier]["structure_stage"] = "final_prediction"
        selected = {self._identifier(row, i) for i, row in enumerate(self._csv("05_final/selected_candidates.csv"))}
        result = []
        for identifier in order:
            row = records[identifier]
            row.setdefault("model_record", False)
            row["candidate_id"] = identifier
            if row.get("error") and row.get("pred_dir"):
                try:
                    from ..screen import prediction_failure_reason
                    reason = prediction_failure_reason(self.artifact_path(row["pred_dir"]))
                    if reason:
                        row["original_error"] = row["error"]
                        row["error"] = reason
                except (ValueError, OSError):
                    pass
            row["evaluated"] = row.get("evaluated", False)
            if comparison and row["evaluated"] and row.get("cif_path"):
                predictor = row.get("predictor_engine") or row.get("engine")
                if predictor:
                    # Independent final predictor structures cannot retain the
                    # generation/refold stage inherited from the source CSV.
                    if row.get("structure_stage") and row["structure_stage"] != predictor:
                        row.setdefault("input_structure_stage", row["structure_stage"])
                    row.update(self._observation_stage({**row, "structure_stage": predictor}))
            row["selected"] = identifier in selected
            row["exclusion_reasons"] = [reason for reason in str(row.get("exclusion_reason") or "").split(";") if reason]
            if not row["evaluated"] and row.get("design_screening_reason"):
                row["exclusion_reasons"].append(row["design_screening_reason"])
            if not row["evaluated"]:
                if not row.get("model_record") and self.config.get("run", {}).get("mode") == "mature":
                    prior = self.state.get("phases", {}).get("mature", {})
                    if identifier == "WT" and self.config.get("rank", {}).get("include_wt_control") is False:
                        row["exclusion_reasons"].append("wt_control_not_in_model_pool")
                    elif prior.get("evaluated") is True:
                        prior_reason = ("sequence_prior_prefilter_unmet" if row.get("pass_llr") is False else
                                        "sequence_prior_top_n_not_selected" if row.get("pass_llr") is True else
                                        "sequence_prior_prefilter_unassessed")
                        row["exclusion_reasons"].append(prior_reason)
                row["exclusion_reasons"] = list(dict.fromkeys(row["exclusion_reasons"] + ["not_scored" if row.get("model_record") else "prefilter_not_selected"]))
                row["final_eligible"] = None
                for key in GEOMETRY_METADATA:
                    if key in row:
                        row[key] = None
            if "final_evaluated" in row:
                # Final metrics belong exclusively to the Boltz cohort. A
                # legacy/imported row cannot turn CF metrics into final values
                # merely by having a representative first-stage structure.
                if not row["evaluated"]:
                    for key in METRICS + ADDITIONAL_METRICS + ("rank", "final_rank", "composite_score"):
                        row[key] = None
                row["first_stage_evaluated"] = bool(
                    not row.get("colabfold_error") and any(
                        row.get("colabfold_" + metric) is not None for metric in METRICS))
            row["missing_metrics"] = [metric for metric in METRICS if row.get(metric) is None]
            for metric in METRICS:
                row.setdefault(metric, None)
            row["quality_eligible"] = self._quality_eligible(row)
            row["quality_eligibility_source"] = (
                "hotspot_only_exclusion_removed" if row["quality_eligible"] is True and row.get("final_eligible") is False
                else "recorded" if row["quality_eligible"] is not None else "not_assessed")
            row.update(self._assessment_details(row))
            _geometry_metadata(row)
            row["metric_provenance"] = metric_provenance(row)
            if row.get("cif_path"):
                try:
                    path = self.artifact_path(row["cif_path"])
                    row["structure_available"] = path.is_file()
                except (OSError, ValueError):
                    row["structure_available"] = False
                if not row["structure_available"]:
                    self.issues.append(f"{identifier}: 구조 파일을 찾지 못했습니다.")
            else:
                row["structure_available"] = False
            result.append(row)
        return [model_observation_record(row, mode=mode)
                for row in self._attach_pose_assessment(result)]

    def _attach_pose_assessment(self, rows):
        """Read persisted assessments only when bound to this exact scored table.

        Loading an old report never computes geometry or silently refreshes a
        candidate's policy. The digest and model identities must both match.
        """
        relative = "03_score/pose_assessment.json"
        if not self._path(relative).exists():
            return rows
        recorded = self._json(relative)
        table = self._path(TABLES[0])
        expected = recorded.get("source_csv_sha256")
        annotations = recorded.get("models")
        identifiers = [self._identifier(row, index) for index, row in enumerate(self._csv(TABLES[0]))]
        valid = (type(recorded.get("schema_version")) is int and recorded["schema_version"] in {1, 2} and isinstance(expected, str)
                 and table.is_file() and expected == sha256(table)
                 and isinstance(annotations, list) and all(isinstance(item, dict) for item in annotations))
        if valid:
            model_ids = [item.get("model_id") for item in annotations]
            valid = (len(set(identifiers)) == len(identifiers) and len(model_ids) == len(identifiers)
                     and all(isinstance(identifier, str) and identifier for identifier in model_ids)
                     and len(set(model_ids)) == len(model_ids) and set(model_ids) == set(identifiers))
        if not valid:
            issue = "자세 평가 기록의 점수표 해시·모델 대응이 일치하지 않아 추가 기록을 적용하지 않았습니다."
            if issue not in self.issues:
                self.issues.append(issue)
            return rows
        by_id = {item["model_id"]: item for item in annotations}
        for row in rows:
            annotation = by_id.get(row["candidate_id"], {})
            row.update({key: copy.deepcopy(annotation[key]) for key in POSE_ANNOTATIONS if key in annotation})
        return rows

    def _quality_eligible(self, row):
        """Remove only a documented hotspot exclusion, without rescoring archives."""
        recorded = row.get("final_eligible")
        if not row["evaluated"] or not isinstance(recorded, bool):
            return None
        if recorded is True or self.config.get("specificity", {}).get("enabled"):
            return recorded
        if set(row["exclusion_reasons"]) != {"hotspot"} or row.get("error") or row["candidate_id"] == "WT":
            return False
        checks = [row.get(key) for key in ("pass_ipsae", "pass_iptm", "pass_plddt", "pass_clash", "dev_cdr_evaluated")]
        if any(value is False for value in checks):
            return False
        if any(value is not True for value in checks):
            return None
        severe = row.get("dev_severe")
        maximum = self.config.get("rank", {}).get("max_dev_severe")
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in (severe, maximum)):
            return None
        return severe <= maximum

    def _assessment_details(self, row):
        """Explain recorded assessments without rescoring read-only results."""
        settings = (self.assessment.get("assessment_config") or {}) if self.assessment else self.config
        thresholds = settings.get("score", {}).get("thresholds", {})
        rank = settings.get("rank", {})
        source_name = (self.assessment["source"] if self.assessment["status"] == "verified"
                       else self.assessment["reason"]) if self.assessment else "run_config"
        recorded_checks = None if self.assessment else row.get("threshold_checks")
        checks = copy.deepcopy(recorded_checks) if isinstance(recorded_checks, list) else []
        for metric, threshold_key, flag, operator in (
                ("ipsae", "ipsae", "pass_ipsae", ">="),
                ("iptm", "iptm", "pass_iptm", ">="),
                ("interface_plddt", "plddt_interface", "pass_plddt", ">="),
                ("hotspot_contact_frac", "hotspot_contact_frac", "pass_hotspot", ">="),
                ("n_clashes", "max_clashes", "pass_clash", "<="),
                ("dev_severe", "max_dev_severe", None, "<=")):
            source = rank if metric == "dev_severe" else thresholds
            recorded = isinstance(source, dict) and threshold_key in source
            if isinstance(recorded_checks, list):
                continue
            passed = row.get(flag) if flag else None
            if self.assessment and self.assessment["status"] != "verified":
                passed = None
            checks.append({"metric": metric, "value": row.get(metric), "operator": operator,
                           "threshold": source[threshold_key] if recorded else None,
                           "passed": passed,
                           "threshold_source": source_name if recorded or self.assessment else "not_recorded"})
        cdr = self.config.get("mature", {}).get("cdr", {})
        if not isinstance(cdr, dict):
            cdr = {}
        mature_state = self.state.get("phases", {}).get("mature", {})
        if not isinstance(mature_state, dict):
            mature_state = {}
        def reasons(key):
            return [value for value in str(row.get(key) or "").split(";") if value]
        return {"threshold_checks": checks,
                **({"assessment_status": self.assessment["status"], "assessment_source": source_name}
                   if self.assessment else {}),
                "cdr_details": {"evaluated": row.get("dev_cdr_evaluated"),
                                "missing_regions": reasons("dev_cdr_missing_regions"),
                                "configured_source": cdr.get("source"), "scheme": cdr.get("scheme"),
                                "configured_ranges": cdr.get("regions"),
                                "recorded_ranges": mature_state.get("cdr_ranges"),
                                "mutation": {key: row[key] for key in ("region", "position", "wt", "mut", "wt_aa", "mut_aa", "mutation") if key in row}},
                "developability_details": {"warnings": reasons("dev_warnings"),
                                            "unevaluated": reasons("dev_unevaluated"),
                                            "severe_count": row.get("dev_severe"),
                                            "unevaluated_count": row.get("dev_severe_unevaluated"),
                                            "flags": {key: value for key, value in row.items()
                                                      if key.startswith("dev_") and isinstance(value, bool)}}}

    def candidates(self, *, search="", sort="composite_score", descending=True, eligible=None,
                   eligibility_basis="recorded", hotspot_min=None, engine_filter="all", offset=0, limit=100):
        if eligibility_basis not in ("recorded", "quality"):
            raise ValueError("적격 판정 기준은 recorded/quality 중 하나여야 합니다.")
        hotspot_min = _hotspot_minimum(hotspot_min)
        rows = self.all_rows()
        total = len(rows)
        if engine_filter != "all":
            rows = [row for row in rows if (row.get("predictor_engine") or row.get("engine")
                    or self.config.get("screen", {}).get("engine")) == engine_filter]
        if search:
            needle = search.lower()
            rows = [row for row in rows if needle in " ".join(
                str(value) for key in SEARCH_FIELDS
                for value in (row[key] if isinstance(row.get(key), list) else [row.get(key)])
                if value is not None).lower()]
        if eligible is not None:
            column = "quality_eligible" if eligibility_basis == "quality" else "final_eligible"
            rows = [row for row in rows if row.get(column) is eligible]
        if hotspot_min is not None:
            rows = [row for row in rows if (fraction := hotspot_filter_fraction(row)) is not None
                    and fraction >= hotspot_min]
        if sort:
            present = [row for row in rows if row.get(sort) is not None]
            missing = [row for row in rows if row.get(sort) is None]
            present.sort(key=lambda row: (isinstance(row[sort], str), row[sort] if isinstance(row[sort], (int, float, str)) else str(row[sort])), reverse=descending)
            rows = present + missing
        columns = list(dict.fromkeys(key for row in rows for key in row))
        return {"total": total, "filtered_total": len(rows), "rows": rows[max(0, offset):max(0, offset) + limit] if limit is not None else rows,
                "columns": columns, "issues": list(dict.fromkeys(self.issues)), "metric_help": METRIC_HELP}

    def summary(self, *, delivery_artifact=False, report_artifact=False):
        if self.retained:
            summary = copy.deepcopy(self.retained["summary"])
            summary.update(retention="minimal", full_execution_available=False,
                           capabilities=self.retained.get("capabilities", {}))
            if self.assessment:
                self.assessment = self._specificity_assessment()
                summary["assessment"] = self.assessment
            return summary
        recorded_engines = self.engine_provenance()
        rows = self.all_rows()
        population = {"candidate_record_total": len(rows),
                      "model_record_total": sum(row.get("model_record") is True for row in rows),
                      "scored_record_total": sum(_assessed(row) for row in rows),
                      "preprediction_record_total": sum(row.get("model_record") is False for row in rows),
                      "population_scope": "Candidate records include prepared/generated/prefiltered sequences. Model records are the authoritative final score/complex ledger, including failed/incomplete attempts; evaluated models are counted separately."}
        from ..publication import publication_issue, tables_available
        final = self.state.get("phases", {}).get("final", {})
        source_delivery_status = final.get("delivery_status")
        source_report_status = final.get("report_status")
        # A successfully written report is its own publication receipt. Its
        # snapshot must not depend on the later run-state write that confirms
        # that same report. Computation/exports failures remain authoritative.
        report_finalized = (report_artifact and source_report_status == "pending"
                            and final.get("exports_status") in {"complete", "not_applicable"}
                            and final.get("tables_committed") is True
                            and tables_available(self.root, final))
        if report_artifact and source_report_status == "pending":
            final = {k: v for k, v in final.items() if k != "report_status"}
        if (delivery_artifact or report_artifact) and source_delivery_status == "pending":
            # The exported artifact is built now. Its source run is still
            # awaiting storage; this exception never authorizes retention.
            final = {k: v for k, v in final.items() if k != "delivery_status"}
        issue = publication_issue(final)
        if final.get("table_sha256") is not None and not tables_available(self.root, final):
            issue = "최신 계산 표의 해시가 일치하지 않습니다."
        issues = list(dict.fromkeys([*self.issues, *([issue] if issue else [])]))
        if self.config.get("run", {}).get("mode") == "complex":
            failures = sum(bool(r.get("error")) for r in rows)
            execution = self.state.get("execution", {}).get("state", "incomplete")
            if report_finalized and execution == "running":
                execution = "complete"
            status = "partial_failure" if (failures or issues) and execution == "complete" else execution
            expected = self.state.get("phases", {}).get("prediction", {}).get("total_models")
            if type(expected) is not int or expected < 0:
                expected = final.get("total_models")
            if type(expected) is not int or expected < 0:
                expected = None
            return {**population, "id": self.root.name, "mode": "complex", "name": self.config.get("run", {}).get("name", self.root.name),
                    "total": len(rows), "scored_total": sum(_assessed(r) for r in rows), "failed_candidates": failures, "issues": issues,
                    "mock": bool(self.config.get("run", {}).get("mock")), "status": status, "complete": status == "complete", "readonly": True,
                    "config": self.config, "state": self.state, "conditions": self.config.get("complex", {}),
                    "source_delivery_status": source_delivery_status, "delivery_artifact": delivery_artifact,
                    "source_report_status": source_report_status, "report_artifact": report_artifact,
                    "expected_total": expected,
                    "engine_provenance": recorded_engines}

        run = self.config.get("run", {})
        warnings = self.state.get("warnings", [])
        execution_state = self.state.get("execution", {}).get("state")
        execution_started = self.state.get("execution", {}).get("started")
        status_applies = bool(self.platform_status) and (execution_started is None or
                            (isinstance(execution_started, (int, float)) and not isinstance(execution_started, bool)
                             and self.platform_status["updated"] >= execution_started))
        authoritative_status = self.platform_status.get("state") if status_applies else execution_state
        if report_finalized and authoritative_status == "running" and execution_state not in {"failed", "interrupted", "cancelled"}:
            authoritative_status = "complete"
        cascade = self._json("03_score/cascade_summary.json") if self._path("03_score/cascade_summary.json").exists() else None
        recorded_cascade = self.state.get("phases", {}).get("cascade", {})
        if cascade is None and recorded_cascade:
            cascade = recorded_cascade
        comparison = self._json("02_comparison/comparison_status.json") if self.config.get("screen", {}).get("comparison", {}).get("enabled") else {}
        stage_failure = bool((cascade or {}).get("status") == "partial_failure"
                             or recorded_cascade.get("status") == "partial_failure"
                             or comparison.get("status") == "partial_failure")
        failures = sum(bool(row.get("error") or row.get("prediction_error")
                            or row.get("prediction_status") == "failed") for row in rows)
        finalized = "final" in self.state.get("phases", {})
        successful_execution = authoritative_status in {None, "complete"}
        status = authoritative_status or ("complete" if finalized else "incomplete")
        if finalized and successful_execution and (failures or issues or stage_failure):
            status = "partial_failure"
        return {**population, "id": self.root.name, "name": run.get("name", self.root.name), "path": str(self.root),
                "mode": run.get("mode"), "mock": bool(run.get("mock")), "total": len(rows),
                "scored_total": sum(_assessed(row) for row in rows),
                "eligible_total": sum(_assessed(row) and row.get("final_eligible") is True for row in rows),
                "quality_eligible_total": sum(_assessed(row) and row.get("quality_eligible") is True for row in rows),
                "selected_total": sum(row["selected"] for row in rows),
                "failed_candidates": failures,
                "zero_eligible": not any(_assessed(row) and row.get("final_eligible") is True for row in rows),
                "complete": finalized and successful_execution and not failures and not issues and not stage_failure, "status": status,
                "execution_state": execution_state, "warnings": warnings,
                "platform_status": {**self.platform_status, "applies": status_applies} if self.platform_status else None,
                "issues": issues, "config": self.config, "state": self.state,
                **({"assessment": self.assessment} if self.assessment else {}),
                "source_delivery_status": source_delivery_status, "delivery_artifact": delivery_artifact,
                "source_report_status": source_report_status, "report_artifact": report_artifact,
                "expected_total": comparison.get("expected_rows"),
                "engine_provenance": recorded_engines,
                "ranking_policy": self._json("03_score/ranking_policy.json") if self._path("03_score/ranking_policy.json").is_file() else None,
                "conditions": {"seed": run.get("seed"), "engine": self.config.get("screen", {}).get("engine"),
                               "cascade": cascade,
                               "comparison": self.config.get("screen", {}).get("comparison", {}),
                               "comparison_progress": comparison,
                               "msa": self.config.get("structure", {}).get("msa", {}),
                               "restraint": self.config.get("screen", {}).get("restraint", {}),
                               "target": self.config.get("target", {})}, "readonly": True}

    def observations(self, candidate_id=None):
        if self.retained:
            rows = [_geometry_metadata(row) for row in copy.deepcopy(self.retained["observations"]["rows"])]
            return {"rows": [r for r in rows if candidate_id is None or r["candidate_id"] == candidate_id],
                    "issues": self.issues}
        if self.config.get("run", {}).get("mode") == "complex":
            return {"rows": [{**r, "observation_id": "complex:"+r["candidate_id"]} for r in self.all_rows() if candidate_id is None or r["candidate_id"] == candidate_id], "issues": self.issues}
        from .. import score_complex
        from ..config import cfg_seed
        rows = []
        candidates = {row["candidate_id"]: row for row in self.all_rows()}
        scored = {identifier: row for identifier, row in candidates.items() if row["evaluated"]}
        from ..utils.predictions import current_manifests
        paths = sorted(current_manifests(self.root, self.issues.append))
        if not paths and self._has_archive_manifest():
            issue = "현재 예측 목록이 없고 보존 기록만 있습니다. 아래 관측은 점수표에서 읽었으며 예측 목록의 모델 수·실행 조건을 복원하지 않습니다."
            if issue not in self.issues:
                self.issues.append(issue)
        for path in paths:
            # Saved attempts remain downloadable audit evidence, but must not
            # duplicate active models or turn a recovered run into a failure.
            # Inspect only descendants: the run's own ancestors may have any name.
            if {".history", ".compact", "source_snapshot"}.intersection(path.relative_to(self.root).parts[:-1]):
                continue
            if "batches" in path.parts and any(parent.joinpath("compact_manifest.json").exists() for parent in path.parents if parent.is_relative_to(self.root)):
                continue
            for index, row in enumerate(self._csv(str(path.relative_to(self.root)))):
                row["candidate_id"] = self._identifier(row, index)
                row["observation_id"] = str(path.relative_to(self.root)) + ":" + str(index)
                recorded_seed = self.config.get("run", {}).get("seed")
                if row.get("engine") in {"boltz", "boltz2"} and (recorded_seed is not None or self.config.get("score", {}).get("boltz2", {}).get("seed") is not None):
                    recorded_seed = cfg_seed(self.config, "score.boltz2.seed")
                row.setdefault("seed_source", "manifest" if "seed" in row else "config")
                row["seed"] = row.get("seed", recorded_seed)
                # Stage manifests describe actual engine conditions. In
                # particular CF's single-sequence/no-restraint settings must
                # not be replaced by the final engine's configuration.
                row.setdefault("msa", row.get("msa_used", self.config.get("structure", {}).get("msa", {}).get("backend")))
                row.setdefault("restraint", self.config.get("screen", {}).get("restraint", {}).get("enabled"))
                row["source_table"] = str(path.relative_to(self.root))
                candidate = candidates.get(row["candidate_id"], {})
                if candidate.get("input_structure_stage"):
                    row.setdefault("input_structure_stage", candidate["input_structure_stage"])
                stage = self._observation_stage(row, row["source_table"])
                row.update(stage)
                # Prediction manifests do not establish geometry for a model.
                # Restore receipts only from a score matching its structure.
                for key in GEOMETRY_METADATA:
                    row[key] = None
                if row.get("structure_pruned"):
                    candidate = candidates.get(row["candidate_id"], {})
                    for metric in METRICS + ADDITIONAL_METRICS:
                        row[metric] = candidate.get("colabfold_" + metric)
                    row.update(structure_available=False, scored=True, score_stage="colabfold")
                    rows.append(row)
                    continue
                structures = []
                try:
                    directory = self.artifact_path(row.get("pred_dir"))
                    if directory.is_dir():
                        structures = sorted(directory.glob("pred.model_idx_*.cif")) or sorted(directory.glob("*_model_*.cif")) or sorted(directory.glob("*.cif")) or sorted(directory.glob("*.pdb"))
                    elif directory.is_file():
                        structures = [directory]
                except (ValueError, OSError) as exc:
                    self.issues.append(f"{row['candidate_id']}: 관측 구조를 읽지 못했습니다 ({exc})")
                for sample, structure_path in enumerate(structures or [None]):
                    observation = dict(row)
                    if sample:
                        observation["observation_id"] += ":model:" + structure_path.name
                    observation["model_index"] = sample
                    if structure_path:
                        observation["cif_path"] = str(structure_path.relative_to(self.root))
                        try:
                            prediction = self._prediction_for_path(structure_path)
                            confidence = score_complex.load_confidence(prediction)
                            if not isinstance(confidence, dict):
                                raise ValueError("신뢰도 데이터는 JSON 객체여야 합니다.")
                            observation.update(native(confidence))
                            observation.update(self._model_identity(confidence, row.get("seed")))
                        except (ValueError, OSError, KeyError) as exc:
                            self.issues.append(f"{observation['observation_id']}: 신뢰도 파일 손상 ({exc})")
                    candidate = candidates.get(row["candidate_id"])
                    # First-stage metrics describe only its representative
                    # model; additional models retain their own confidence.
                    first_path = candidate.get("colabfold_cif_path") if candidate else None
                    if candidate and structure_path and first_path:
                        try:
                            if self.artifact_path(first_path) == structure_path.resolve():
                                for metric in METRICS + ADDITIONAL_METRICS:
                                    if metric in GEOMETRY_METADATA or candidate.get("colabfold_" + metric) is not None:
                                        observation[metric] = candidate.get("colabfold_" + metric)
                                observation["scored"] = True
                                observation["score_stage"] = "colabfold"
                                for key in ("iptm_scope", "iptm_source"):
                                    if "colabfold_" + key in candidate:
                                        observation[key] = candidate["colabfold_" + key]
                        except (ValueError, OSError):
                            pass
                    if candidate and observation["stage"] == "first" and candidate.get("colabfold_error"):
                        observation.setdefault("error", candidate["colabfold_error"])
                    final = scored.get(row["candidate_id"])
                    if final and structure_path and final.get("cif_path"):
                        try:
                            if self.artifact_path(final["cif_path"]) == structure_path.resolve():
                                for metric in METRICS + ADDITIONAL_METRICS + ("composite_score",):
                                    if metric in GEOMETRY_METADATA or final.get(metric) is not None:
                                        observation[metric] = final.get(metric)
                                observation["scored"] = True
                                observation["score_stage"] = "boltz2" if observation["stage"] == "final" else observation["structure_stage"]
                                for key in ("iptm_scope", "iptm_source"):
                                    if key in final:
                                        observation[key] = final[key]
                        except (ValueError, OSError):
                            pass
                    rows.append(observation)
        # Normal runs may only have scored rows; preserve each score's own condition.
        if not rows:
            for row in candidates.values():
                if not row["evaluated"] and not row.get("structure_available"):
                    continue
                observation = {**row, "observation_id": "scored:" + row["candidate_id"]}
                observation.update(source_table="03_score/scored_candidates.csv",
                                   observation_source="scored_table_only", prediction_manifest_status="not_recorded")
                observation.update(self._observation_stage(row))
                observation.update(self._row_conditions(row))
                if observation["stage"] == "first":
                    for metric in METRICS + ADDITIONAL_METRICS:
                        observation[metric] = row.get("colabfold_" + metric)
                if row.get("cif_path"):
                    try:
                        confidence = score_complex.load_confidence(self._prediction_for_path(self.artifact_path(row["cif_path"])))
                        if isinstance(confidence, dict):
                            observation.update(self._model_identity(confidence, observation.get("seed")))
                    except (ValueError, OSError, KeyError) as exc:
                        self.issues.append(f"{observation['observation_id']}: 신뢰도 파일 손상 ({exc})")
                rows.append(observation)
        if candidate_id:
            rows = [row for row in rows if row["candidate_id"] == candidate_id]
        for row in rows:
            _geometry_metadata(row)
            row["metric_provenance"] = metric_provenance(row)
        return {"rows": rows, "issues": list(dict.fromkeys(self.issues))}

    def _has_archive_manifest(self):
        """Detect ambiguous archive-only records without reading their contents."""
        from ..utils.predictions import ARCHIVE_DIRECTORIES
        for directory, children, files in os.walk(self.root, followlinks=False):
            parent = Path(directory)
            children[:] = [name for name in children if not (parent / name).is_symlink()]
            if "predictions.csv" in files and any(part in ARCHIVE_DIRECTORIES for part in parent.relative_to(self.root).parts):
                if not (parent / "predictions.csv").is_symlink():
                    return True
        return False

    def _observation_stage(self, row, source_table=""):
        cascade = bool(self.config.get("screen", {}).get("cascade", {}).get("enabled")
                       or self.state.get("phases", {}).get("cascade"))
        engine = row.get("engine")
        structure_stage = row.get("structure_stage")
        if self.config.get("screen", {}).get("comparison", {}).get("enabled") is True or row.get("comparison") is True:
            predictor = row.get("predictor_engine") or engine or structure_stage or "unknown"
            name = {"colabfold": "ColabFold", "boltz2": "Boltz-2", "boltz": "Boltz-2",
                    "chai1": "Chai-1", "opendde": "OpenDDE Preview"}.get(predictor, str(predictor))
            stage = {"stage": "prediction", "structure_stage": predictor,
                     "stage_label": name + " 독립 예측"}
            if row.get("input_structure_stage"):
                stage["input_structure_stage"] = row["input_structure_stage"]
            return stage
        if engine == "colabfold" or structure_stage == "colabfold" or cascade and source_table.startswith(("02_screen/", "02b_mature_complex/")):
            return {"stage": "first", "structure_stage": "colabfold", "stage_label": "ColabFold 1차"}
        if cascade and (engine in {"boltz", "boltz2"} or source_table.startswith("02c_boltz2/") or structure_stage == "boltz2"):
            return {"stage": "final", "structure_stage": "boltz2", "stage_label": "Boltz-2 최종"}
        return {"stage": "prediction", "structure_stage": structure_stage or engine or "unknown", "stage_label": str(engine or "예측")}

    def _row_conditions(self, row):
        from ..config import cfg_seed
        prefix = "colabfold_" if self._observation_stage(row)["stage"] == "first" else "boltz_"
        engine = row.get("engine")
        seed = cfg_seed(self.config, "score.boltz2.seed") if engine in {"boltz", "boltz2"} else self.config.get("run", {}).get("seed")
        return {"seed": row.get("seed", row.get(prefix + "seed", seed)),
                "msa": row.get("msa", row.get(prefix + "msa", self.config.get("structure", {}).get("msa", {}).get("backend"))),
                "restraint": row.get("restraint", row.get(prefix + "restraint", self.config.get("screen", {}).get("restraint", {}).get("enabled")))}

    @staticmethod
    def _model_identity(confidence, base_seed):
        raw_model = str(confidence.get("raw_model") or "")
        seed_match = re.search(r"_seed_(\d+)(?:\.pdb)?$", raw_model)
        model_match = re.search(r"_model_(\d+)_seed_", raw_model)
        result = {}
        if seed_match:
            result.update(base_seed=base_seed, seed=int(seed_match.group(1)), seed_source="raw_model")
        if model_match:
            result["engine_model_id"] = int(model_match.group(1))
        return result

    def _prediction_for_path(self, path):
        from ..score_complex import PredictionFiles
        path = contained(self.root, path, exists=True)
        if path.name.startswith("pred.model_idx_"):
            idx = path.stem.split("_")[-1]
            pae = path.parent / f"scores_extra_idx_{idx}.npz"
            if not pae.is_file():
                pae = path.parent / "scores_extra.npz" if idx == "0" else None
            prediction = PredictionFiles(path, scores_npz=path.parent / f"scores.model_idx_{idx}.npz",
                                         pae_npz=pae, source="chai1")
        elif path.suffix.lower() == ".cif" or path.name.startswith("colabfold_model_"):
            prediction = PredictionFiles(path, confidence_json=path.parent / f"confidence_{path.stem}.json",
                                         pae_npz=path.parent / f"pae_{path.stem}.npz",
                                         source="opendde" if path.name.startswith("opendde_model_") else "colabfold" if path.name.startswith("colabfold_model_") else "boltz2")
        else:
            prediction = PredictionFiles(path)
        for associated in (prediction.confidence_json, prediction.scores_npz, prediction.pae_npz):
            if associated:
                contained(self.root, associated)
        return prediction

    def artifact_path(self, value):
        resolved = resolve_run_artifact(self.root, value)
        if resolved is None:
            raise FileNotFoundError("구조 경로가 없습니다.")
        return contained(self.root, resolved)

    def structure(self, candidate_id, observation_id=None):
        if self.retained:
            from .retention import retained_structure
            return self._coordinate_contract(retained_structure(self.root, candidate_id, observation_id))
        if self.config.get("run", {}).get("mode") == "complex":
            row = next((r for r in self.all_rows() if r["candidate_id"] == candidate_id), None)
            if row is None or not row.get("structure_file"):
                raise FileNotFoundError("이 모델의 구조가 없습니다.")
            data = self._json(row["structure_file"])
            if not data or not isinstance(data.get("content"), str) or not data["content"]:
                raise ValueError("저장된 복합체 구조가 손상되었습니다.")
            if data.get("pae") is not None:
                from ..score_complex import validate_pae
                import numpy as np
                try:
                    pae = data["pae"]
                    labels = pae["labels"]
                    if not isinstance(labels, list) or not all(isinstance(label, str) for label in labels):
                        raise ValueError("invalid_pae_axis")
                    chains = data.get("chains")
                    # Legacy displays may record only chain IDs. Compare axes
                    # only when every chain's residue list was actually saved;
                    # absent metadata is not evidence of a mismatching PAE.
                    if (isinstance(chains, list) and chains
                            and all(isinstance(chain, dict) and isinstance(chain.get("residues"), list)
                                    for chain in chains)):
                        coordinate_labels = [residue["label"] for chain in chains
                                             for residue in chain["residues"]]
                        if not all(isinstance(label, str) for label in coordinate_labels) or labels != coordinate_labels:
                            raise ValueError("invalid_pae_coordinate_axis")
                    validate_pae(np.asarray(pae["values"], dtype=float), len(labels))
                except (ValueError, TypeError, KeyError, OverflowError):
                    # A damaged error matrix cannot invalidate independently
                    # usable saved coordinates or silently remove their export.
                    message = "PAE 자료가 손상되어 표시할 수 없습니다. 표와 구조는 계속 확인할 수 있습니다."
                    self.issues.append(f"{candidate_id}: {message}")
                    data = {**data, "pae": None, "pae_status": "invalid", "pae_error": message}
                    prior = data.get("issues") if isinstance(data.get("issues"), list) else []
                    data["issues"] = list(dict.fromkeys([*prior, message]))
            return self._coordinate_contract({**data, "candidate_id": candidate_id, "observation_id": observation_id}, row=row)
        return self._nanobody_structure(candidate_id, observation_id)

    def _native_coordinate_receipt(self, path, digest, row):
        """Match this exact file to a saved native completion output hash.

        A file name, score-table reference, or parseable coordinate is not a
        native-origin receipt. Search only the selected file's owned ancestors.
        """
        if self.config.get("run", {}).get("mock") or row.get("engine") == "mock":
            return {"status": "unavailable", "reason": "mock_coordinates"}
        markers = (("boltz_complete.json", "boltz_out"),
                   ("colabfold_complete.json", "colabfold_out"),
                   ("prediction_complete.json", "chai_out"),
                   ("prediction_complete.json", "opendde_out"))
        mismatch = False
        for parent in path.parents:
            if not parent.is_relative_to(self.root):
                break
            for name, output in markers:
                marker = parent / name
                if not marker.is_file():
                    continue
                relative = str(marker.relative_to(self.root))
                try:
                    receipt = self._json(relative)
                    directory = self._path(str((parent / output).relative_to(self.root)))
                except (ValueError, OSError):
                    # A stale/external receipt cannot hide usable local display
                    # coordinates or establish their native origin.
                    continue
                outputs = receipt.get("outputs")
                if receipt.get("status") != "complete" or not isinstance(outputs, dict):
                    continue
                if not path.is_relative_to(directory):
                    continue
                expected = outputs.get(str(path.relative_to(directory)))
                if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-f]{64}", expected):
                    continue
                if expected != digest:
                    mismatch = True
                    continue
                return {"status": "verified", "basis": "native_completion_output_sha256",
                        "receipt_path": relative, "receipt_sha256": sha256(marker)}
        return {"status": "hash_mismatch" if mismatch else "unavailable",
                "reason": "native_coordinate_hash_mismatch" if mismatch else "native_origin_not_recorded"}

    def _coordinate_contract(self, payload, *, path=None, row=None):
        """Normalize provenance without changing display coordinates or axes.

        UTF-8 bytes are preserved (including CRLF). Embedded raw coordinates
        need their own recorded SHA; legacy display content never establishes
        native origin merely because it resembles an engine output.
        """
        data = dict(payload)
        source = data.get("coordinate_source")
        source = source if isinstance(source, dict) else {}
        row = row or {}
        content = data.get("content", "")
        display_sha = hashlib.sha256(content.encode("utf-8")).hexdigest()
        raw, raw_path, raw_format = None, None, None
        verification = {"status": "unavailable", "reason": "native_origin_not_recorded"}
        expected = source.get("raw_sha256")
        if isinstance(expected, str) and re.fullmatch(r"[0-9a-f]{64}", expected):
            # A saved complex/retained payload explicitly records native bytes.
            embedded = data.get("raw_content")
            raw = embedded.encode("utf-8") if isinstance(embedded, str) else None
            try:
                if source.get("raw_path"):
                    original = self.artifact_path(source["raw_path"])
                    raw_path = str(original.relative_to(self.root))
                    if original.is_file():
                        original_bytes = original.read_bytes()
                        if raw is not None and raw != original_bytes:
                            raise ValueError("embedded_raw_mismatch")
                        raw = original_bytes
                        raw_format = original.suffix.lower().lstrip(".")
                artifacts = row.get("artifacts", {})
                if isinstance(artifacts, dict) and raw_path in artifacts and artifacts[raw_path] != expected:
                    raise ValueError("recorded_raw_hash_conflict")
                if raw is None:
                    verification["reason"] = "recorded_native_coordinates_missing"
                elif hashlib.sha256(raw).hexdigest() != expected:
                    raise ValueError("recorded_raw_hash_mismatch")
                else:
                    verification = {"status": "verified", "basis": "stored_coordinate_source_sha256"}
                    if raw_path and (self.root / raw_path).is_file():
                        verification["basis"] = "stored_coordinate_source_and_file_sha256"
                    if row.get("structure_file"):
                        verification["receipt_path"] = row["structure_file"]
                        verification["receipt_sha256"] = sha256(self.artifact_path(row["structure_file"]))
                    elif data.get("content_file"):
                        verification["receipt_path"] = data["content_file"]
            except (ValueError, TypeError, OSError):
                verification = {"status": "hash_mismatch", "reason": "recorded_native_coordinates_unverified"}
        elif path is not None:
            raw = path.read_bytes()
            expected = hashlib.sha256(raw).hexdigest()
            verification = self._native_coordinate_receipt(path, expected, row)
            raw_path, raw_format = str(path.relative_to(self.root)), path.suffix.lower().lstrip(".")
        verified = verification["status"] == "verified"
        if verified:
            try:
                raw_content = raw.decode("utf-8")
            except UnicodeError:
                verified = False
                verification = {"status": "unavailable", "reason": "native_coordinates_not_utf8"}
        transforms = []
        relation = "unverified_legacy"
        transform = "native origin not verified"
        if verified:
            relation = "native_identical" if raw_content == content else "transformed_display"
            transform = "none" if relation == "native_identical" else "display differs from verified native bytes; transform not recorded"
            if relation == "transformed_display":
                if source.get("display_transform"):
                    transform = source["display_transform"]
                    transforms.append("recorded_display_transform")
                # A VHH reference is a separate overlay. Only the complex
                # writer's successful reference receipt records an applied fit.
                if self.config.get("run", {}).get("mode") == "complex":
                    transforms.append("coordinate_serialization_and_numbering_mapping")
                    aligned = isinstance(data.get("reference"), dict) and data["reference"].get("available") is True
                    if aligned:
                        transforms.append("recorded_reference_alignment")
                    transform = "coordinate serialization and recorded numbering mapping" + ("; recorded reference alignment" if aligned else "")
            raw_format = raw_format or data.get("raw_format")
        data.update(raw_content=raw_content if verified else None, raw_format=raw_format if verified else None,
                    coordinate_source={"schema_version": 2, "raw_preserved": verified,
                        "raw_path": raw_path, "raw_sha256": expected if verified else None,
                        "display_sha256": display_sha, "display_relation": relation,
                        "display_transform": transform, "display_transforms": transforms,
                        "verification": verification})
        from ..result_schema import structure_payload
        return structure_payload(data, observation=row,
                                 mode=self.config.get("run", {}).get("mode", "unknown"))

    def _nanobody_structure(self, candidate_id, observation_id=None):
        from .structure_review import closest_atoms, cysteine_geometry, cdr_annotations, reference_overlay
        from .. import score_complex
        from ..utils import pdb_io
        rows = self.all_rows()
        row = next((row for row in rows if row["candidate_id"] == candidate_id), None)
        if row is None:
            raise KeyError("후보가 없습니다.")
        if observation_id:
            observation = next((obs for obs in self.observations(candidate_id)["rows"] if obs["observation_id"] == observation_id), None)
            if observation is None:
                raise KeyError("관측이 없습니다.")
            row = {**row, **observation}
            # A failed/missing first-stage observation must not borrow the
            # candidate's successful final representative structure.
            row["cif_path"] = None
            if observation.get("cif_path"):
                row["cif_path"] = str(self.artifact_path(observation["cif_path"]))
            elif observation.get("pred_dir"):
                directory = self.artifact_path(observation["pred_dir"])
                row["cif_path"] = str(score_complex.find_prediction_files(directory).cif)
        path = self.artifact_path(row.get("cif_path"))
        if not path.is_file():
            raise FileNotFoundError("이 후보의 구조 파일이 없습니다.")
        structure = pdb_io.load_structure(path)
        model = next(structure.get_models(), None)
        if model is None or next(model.get_atoms(), None) is None:
            raise ValueError("이 후보의 구조를 표시할 수 없습니다: 구조 파일에 원자 좌표가 없습니다.")
        labels, index_map = score_complex.build_index_map(structure)
        chains = []
        target_chain, binder_chain = str(row.get("target_chain") or "A"), str(row.get("binder_chain") or "B")
        for chain in model:
            chains.append({"id": chain.id, "role": "target" if chain.id == target_chain else "binder" if chain.id == binder_chain else "context",
                           "residues": [{"label": pdb_io.residue_label(res, chain.id), "number": res.id[1], "icode": res.id[2].strip(), "name": res.resname,
                                         "ca": pdb_io.ca_coord(res).tolist() if pdb_io.ca_coord(res) is not None else None,
                                         "atom_ids": {atom.name: atom.get_serial_number() for atom in res.get_atoms()}}
                                        for res in pdb_io.iter_residues(chain)]})
        contacts, clashes = [], []
        binder = pdb_io.get_chain(structure, binder_chain)
        contact_cutoff = float(self.config.get("score", {}).get("contact_cutoff", 5.0))
        clash_cutoff = float(self.config.get("score", {}).get("clash_cutoff", 2.0))
        for chain in model:
            if chain.id == binder_chain:
                continue
            for pair in pdb_io.contact_pairs(chain, binder, cutoff=max(contact_cutoff, clash_cutoff), atom_mode="heavy"):
                item = {"a": pair["a"], "b": pair["b"], "distance": float(pair["dist"]), "kind": "target" if chain.id == target_chain else "context"}
                if item["distance"] <= contact_cutoff:
                    contacts.append(item)
                if item["distance"] < clash_cutoff:
                    clashes.append({**item, **closest_atoms(pair), "cutoff": clash_cutoff,
                                    "distance_below_cutoff": clash_cutoff - item["distance"]})
        target_state = dict(self.state.get("phases", {}).get("target", {}))
        hotspots = []
        if target_state:
            try:
                from ..cli import _prediction_hotspot_labels
                if target_state.get("structure"):
                    target_state["structure"] = str(self.artifact_path(target_state["structure"]))
                hotspots = _prediction_hotspot_labels(path, target_chain, target_state)
            except (ValueError, OSError, KeyError) as exc:
                self.issues.append(f"hotspot 번호 대응 미평가: {exc}")
        if not hotspots and not target_state:
            requested = self.config.get("target", {}).get("hotspots", [])
            if isinstance(requested, list):
                try:
                    hotspots = score_complex.resolve_hotspot_labels(requested, target_chain, index_map)
                except (ValueError, TypeError) as exc:
                    self.issues.append(f"hotspot 번호 대응 미평가: {exc}")
        pae = None
        try:
            prediction = self._prediction_for_path(path)
            matrix = score_complex.load_pae(prediction)
            if matrix is not None:
                score_complex.validate_pae(matrix, len(labels))
                pae = {"values": matrix.tolist(), "labels": labels, "unit": "Å"}
        except (ValueError, OSError, KeyError) as exc:
            self.issues.append(f"PAE 미평가 또는 손상: {exc}")
        conditions = self._row_conditions(row)
        try:
            confidence = score_complex.load_confidence(self._prediction_for_path(path))
            if isinstance(confidence, dict):
                conditions.update(self._model_identity(confidence, conditions.get("seed")))
        except (ValueError, OSError, KeyError):
            # The observation reader separately reports corrupt confidence;
            # keep the structure and its independently validated PAE usable.
            pass
        mapping = self.residue_mapping()
        payload = native({"candidate_id": candidate_id, "observation_id": observation_id, "content": path.read_text(),
                       "engine": row.get("engine"), **self._observation_stage(row), **conditions,
                       "format": "pdb" if path.suffix.lower() == ".pdb" else "cif", "chains": chains,
                       "target_chain": target_chain, "binder_chain": binder_chain, "hotspots": hotspots,
                       "contacts": contacts, "clashes": clashes, "pae": pae,
                       "residue_mapping": mapping, "cdr": cdr_annotations(self, row, binder),
                       "cysteines": cysteine_geometry(binder), "reference": reference_overlay(self, pdb_io.get_chain(structure, target_chain), mapping),
                       "target_region": self.config.get("input_review", {}).get("target_region"),
                       "issues": list(dict.fromkeys(self.issues))})
        return self._coordinate_contract(payload, path=path, row=row)

    def residue_mapping(self):
        from ..utils.residue_mapping import read_residue_mapping
        numeric_fields = {"output_position", "reference_position", "source_full_position",
                          "source_observed", "source_standard_observed", "hotspot_excluded"}
        for relative in ("01_target/residue_mapping.csv", "residue_mapping.csv", "01_target/target_mapping.csv"):
            path = self._path(relative)
            if path.exists():
                try:
                    # Author labels are opaque text, including negative numbers,
                    # leading zeros and literal apostrophes. Only the explicitly
                    # typed output-position/assessment fields are normalized.
                    return [{key: _cell(value, key) if key in numeric_fields else value
                             for key, value in row.items()} for row in read_residue_mapping(path)]
                except (ValueError, csv.Error, UnicodeError, OSError) as exc:
                    self.issues.append(f"{relative}: 읽을 수 없는 CSV ({exc})")
                    return []
        return []

    def artifacts(self):
        result = []
        for path in sorted(self.root.rglob("*")):
            if path.is_file():
                try:
                    contained(self.root, path)
                except ValueError:
                    continue
                relative = str(path.relative_to(self.root))
                if path.name.startswith(".run.lock"):
                    continue
                result.append({"path": relative, "size": path.stat().st_size, "type": path.suffix.lower().lstrip(".")})
        return {"artifacts": result, "provenance": self.provenance()}

    def provenance(self):
        authoritative = []
        for relative in ("run_config.yaml", "run_state.json", "complex_results.json", "03_score/scored_candidates.csv",
                         "05_final/selected_candidates.csv", "03_score/ranking_policy.json", "retained/manifest.json",
                         "02_comparison/comparison_identity.json", "02_comparison/comparison_status.json", "preflight.json",
                         "06_specificity/assessment.json"):
            path = self._path(relative)
            if path.is_file():
                path = self.artifact_path(relative)
                authoritative.append({"path": relative, "size": path.stat().st_size, "sha256": sha256(path)})
        return {"source_artifacts": authoritative,
            "engine_provenance": self.engine_provenance(),
            "ranking_policy": self._json("03_score/ranking_policy.json") if self._path("03_score/ranking_policy.json").is_file() else None,
            "scope": "saved authoritative result files; raw predictions listed separately"}

    def engine_provenance(self):
        """Expose saved versions and runtime fingerprints without a live probe."""
        identity = self._json("02_comparison/comparison_identity.json")
        runtime = identity.get("runtime", {}) if isinstance(identity, dict) else {}
        preflight = self._json("preflight.json")
        registered = preflight.get("engines", {}) if isinstance(preflight, dict) else {}
        if not isinstance(runtime, dict):
            runtime = {}
        if not isinstance(registered, dict):
            registered = {}
        engines = list(dict.fromkeys([*runtime, *registered]))
        records = []
        for engine in engines:
            saved = registered.get(engine, {})
            saved = saved if isinstance(saved, dict) else {}
            probe = saved.get("env_probe", {})
            probe = probe if isinstance(probe, dict) else {}
            versions = probe.get("imports", {})
            versions = versions if isinstance(versions, dict) else {}
            versions = {name: value.get("version") for name, value in versions.items()
                        if isinstance(value, dict) and isinstance(value.get("version"), str)}
            native_runtime = runtime.get(engine)
            native_runtime = copy.deepcopy(native_runtime) if isinstance(native_runtime, dict) else None
            if native_runtime and isinstance(native_runtime.get("source"), dict):
                # Paths as JSON object keys cannot be redacted without losing
                # colliding external files. Preserve every digest in path rows
                # so the portable presentation copy can sanitize each value.
                native_runtime["source_files"] = [{"path": path, "sha256": digest}
                                                  for path, digest in native_runtime.pop("source").items()]
            records.append({"engine": engine, "versions": versions,
                            "runtime": native_runtime,
                            "weights": saved.get("weights"), "installation_identity": saved.get("installation_identity"),
                            "version_source": "preflight.json" if versions else "not_recorded",
                            "runtime_source": "02_comparison/comparison_identity.json" if native_runtime else "not_recorded"})
        return records

    def _fasta(self, rows):
        """Use the same complete sequence set for API and offline downloads."""
        if self.config.get("run", {}).get("mode") == "complex":
            from ..complex import proteins
            entries = proteins(self.config["complex"]["proteins"], historical=True)
            return "".join(f">{p['id']} {' '.join(p['name'].split())} original_positions={','.join(map(str, p['original_positions']))}\n{p['sequence']}\n" for p in entries)
        return "".join(f">{row['candidate_id']}\n{row['sequence']}\n" for row in rows if row.get("sequence"))

    def export_to_file(self, format: str, filters: dict | None = None):
        """Spool a full archive for HTTP delivery; caller removes the temp file.

        The legacy bytes interface is retained for programmatic callers. The
        public full-archive API uses this method so raw multi-GB ZIP bytes never
        accumulate in a BytesIO or Response body.
        """
        if format != "zip":
            raise ValueError("파일 스트리밍은 전체 원본 ZIP에 사용합니다.")
        descriptor, name = tempfile.mkstemp(prefix="nanocofold-export-", suffix=".zip")
        os.close(descriptor)
        path = Path(name)
        try:
            _, media_type, filename = self.export(format, filters, destination=path)
            return path, media_type, filename
        except BaseException:
            path.unlink(missing_ok=True)
            raise

    def export(self, format: str, filters: dict | None = None, *, destination: Path | None = None, report_artifact=False):
        if report_artifact and format != "html":
            raise ValueError("보고서 발행 스냅숏은 HTML 형식에만 적용합니다.")
        allowed = _exploration({} if filters is None else filters)
        query = {key: value for key, value in allowed.items() if key in {"search", "sort", "descending", "eligible", "eligibility_basis", "hotspot_min", "engine_filter"}}
        eligibility = allowed.get("eligibility")
        if eligibility in {"all", "missing"}:
            query.pop("eligible", None)
        elif eligibility in {"true", "false"}:
            query["eligible"] = eligibility == "true"
        table = self.candidates(**query, limit=None)
        if eligibility == "missing":
            column = "quality_eligible" if allowed.get("eligibility_basis") == "quality" else "final_eligible"
            table["rows"] = [row for row in table["rows"] if row.get(column) is None]
            table["filtered_total"] = len(table["rows"])
        # Export the entire candidate set and preserve exploration criteria independently.
        all_table = self.candidates(limit=None)
        rows, columns = all_table["rows"], all_table["columns"]
        if format == "report_zip":
            return self._delivery_export(all_table, allowed, table)
        from ..utils.presentation import portable_metadata
        all_table = portable_metadata(all_table, root=self.root)
        rows = all_table["rows"]
        provenance = {"summary": self.summary(report_artifact=report_artifact), "filters": allowed, "filtered_ids": [row["candidate_id"] for row in table["rows"]],
                      **self.provenance()}
        provenance = portable_metadata(provenance, root=self.root)
        def csv_bytes():
            from ..utils.spreadsheet import csv_cell, csv_headers
            stream = io.StringIO()
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writerow(dict(zip(columns, csv_headers(columns))))
            writer.writerows([{key: csv_cell(json.dumps(native(value), ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value) for key, value in row.items()} for row in rows])
            return ("\ufeff" + stream.getvalue()).encode()
        if format == "csv":
            return csv_bytes(), "text/csv; charset=utf-8", "candidates.csv"
        if format == "fasta":
            return self._fasta(rows).encode(), "text/plain; charset=utf-8", "candidates.fasta"
        if format == "xlsx":
            from openpyxl import Workbook
            book = Workbook()
            sheet = book.active
            sheet.title = "All candidates"
            sheet.append(columns)
            for row in rows:
                values = [json.dumps(native(row.get(key)), ensure_ascii=False, allow_nan=False) if isinstance(row.get(key), (dict, list)) else row.get(key) for key in columns]
                sheet.append(values)
                # Preserve exact identifiers while forcing uploaded text to data,
                # rather than spreadsheet formulas.
                for cell, value in zip(sheet[sheet.max_row], values):
                    if isinstance(value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            manifest = book.create_sheet("Provenance")
            manifest.append(["field", "part", "json"])
            for key, value in provenance.items():
                encoded = json.dumps(native(value), ensure_ascii=False)
                # Excel truncates strings at 32,767 characters. Preserve long
                # config/artifact manifests in ordered, reconstructable chunks.
                from ..utils.spreadsheet import excel_text_chunks
                chunks = excel_text_chunks(encoded)
                for index, chunk in enumerate(chunks):
                    manifest.append([key, f"{index + 1}/{len(chunks)}", chunk])
            from ..utils.spreadsheet import deterministic_xlsx, force_text
            force_text(book)
            return deterministic_xlsx(book), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "candidates.xlsx"
        if format == "html":
            payload = self._offline_payload(all_table, allowed, provenance, report_artifact=report_artifact)
            return self._offline_document(payload), "text/html; charset=utf-8", "report_offline.html"
        if format == "zip":
            stream = io.BytesIO() if destination is None else destination
            with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("candidates.csv", csv_bytes())
                archive.writestr("manifest.json", json.dumps(native(provenance), ensure_ascii=False, indent=2))
                archive.writestr("candidates.fasta", self.export("fasta")[0])
                archive.writestr("report_offline.html", self._offline_document(self._offline_payload(all_table, allowed, provenance)))
                for artifact in self.artifacts()["artifacts"]:
                    archive.write(self.artifact_path(artifact["path"]), "run/" + artifact["path"])
            return stream.getvalue() if destination is None else destination, "application/zip", "result_package.zip"
        raise ValueError("내보내기 형식은 report_zip/csv/xlsx/fasta/html/zip 중 하나입니다.")

    def _offline_payload(self, all_table, filters, provenance, *, delivery_artifact=False, report_artifact=False):
        structures, pae_arrays = [], {}
        for row in all_table["rows"]:
            if row["structure_available"]:
                try:
                    observation = "complex:" + row["candidate_id"] if self.config.get("run", {}).get("mode") == "complex" else None
                    structures.append(_pack_offline_pae(self.structure(row["candidate_id"], observation), pae_arrays))
                except (OSError, ValueError, KeyError) as exc:
                    self.issues.append(f"{row['candidate_id']}: {exc}")
        observations = self.observations()
        for observation in observations["rows"]:
            if self.config.get("run", {}).get("mode") == "complex" or observation.get("structure_pruned"):
                continue
            try:
                structures.append(_pack_offline_pae(self.structure(observation["candidate_id"], observation["observation_id"]), pae_arrays))
            except (OSError, ValueError, KeyError) as exc:
                self.issues.append(f"{observation['observation_id']}: {exc}")
        # Complex model rows contain scores, not the input protein sequences.
        # Expand copies and analyzed ranges once with the same authoritative
        # serializer used by API FASTA and the shareable ZIP.
        try:
            sequences_fasta = self._fasta(all_table["rows"])
        except (KeyError, TypeError, ValueError):
            # A historical/malformed config must not prevent viewing saved
            # scores and structures, or fabricate an empty sequence download.
            sequences_fasta = None
        from ..utils.presentation import portable_metadata
        return portable_metadata({"summary": self.summary(delivery_artifact=delivery_artifact, report_artifact=report_artifact), "candidates": all_table, "observations": observations,
                                  "sequences_fasta": sequences_fasta,
                                  "structures": structures, "pae_arrays": pae_arrays,
                                  "filters": filters, "provenance": provenance}, root=self.root)

    @staticmethod
    def _offline_document(payload, assets=None):
        from ..report import _resolve_viewer_asset
        template = Path(__file__).resolve().parents[1] / "templates/platform_offline.html"
        encoded = json.dumps(native(payload), ensure_ascii=False, allow_nan=False).replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
        if not template.exists():
            return _offline_fallback(encoded).encode()
        document = template.read_text()
        for token, filename in (("__PLOTLY_JS__", "plotly.min.js"), ("__3DMOL_JS__", "3dmol-min.js"),
                                ("__MOLSTAR_JS__", "molstar.js"), ("__MOLSTAR_CSS__", "molstar.css")):
            if assets is None:
                raw, evidence = _resolve_viewer_asset({}, filename, asset_root=ASSET_ROOT, local_only=True)
                if not evidence["ok"]:
                    raise OSError(evidence["blocker"])
            else:
                raw = assets[filename]
            content = raw.decode("utf-8")
            document = document.replace(token, re.sub(r"</(script|style)", r"<\\/\1", content, flags=re.I))
        # Inject records last: literal asset placeholders in names/errors must
        # remain data rather than being replaced by vendor source code.
        return document.replace("__RESULT_DATA__", encoded, 1).encode()

    def _delivery_export(self, all_table, filters, filtered_table):
        import hashlib
        from ..report_delivery import package_report
        from ..report import _resolve_viewer_asset
        assets, asset_sources = {}, {}
        for name in ("molstar.js", "molstar.css", "plotly.min.js", "3dmol-min.js"):
            content, evidence = _resolve_viewer_asset({}, name, asset_root=ASSET_ROOT, local_only=True)
            if not evidence["ok"]:
                raise OSError(evidence["blocker"])
            assets[name], asset_sources[name] = content, evidence
        # Only these source receipts are hashed. Large MSA databases, caches,
        # logs, old archives and intermediate tensors are never traversed here.
        source_paths = ("run_config.yaml", "run_state.json", "complex_results.json", "models.csv",
                        *TABLES, "05_final/selected_candidates.csv", "03_score/ranking_policy.json", "01_target/residue_mapping.csv",
                        "retained/manifest.json", "retained/record.json", "preflight.json",
                        "02_comparison/comparison_identity.json", "02_comparison/comparison_status.json",
                        "06_specificity/assessment.json")
        sources = []
        for relative in source_paths:
            path = self._path(relative)
            if path.is_file():
                path = self.artifact_path(relative)
                sources.append({"path": relative, "size": path.stat().st_size, "sha256": sha256(path)})
        provenance = {"summary": self.summary(), "filters": filters,
                      "filtered_ids": [row["candidate_id"] for row in filtered_table["rows"]],
                      "source_artifacts": sources, "viewer_asset_sources": asset_sources,
                      "engine_provenance": self.engine_provenance(),
                      "ranking_policy": self._json("03_score/ranking_policy.json") if self._path("03_score/ranking_policy.json").is_file() else None,
                      "scope": "report_inputs_and_embedded_structures"}
        payload = self._offline_payload(all_table, filters, provenance, delivery_artifact=True)
        # Hash exactly what the collaborator can inspect, including every seed
        # and PAE. This is not a manifest of the omitted full run directory.
        embedded = []
        for structure in payload["structures"]:
            content = structure["content"].encode()
            entry = {"candidate_id": structure["candidate_id"], "observation_id": structure.get("observation_id"),
                     "format": structure["format"], "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
            if structure.get("pae"):
                entry["pae_sha256"] = hashlib.sha256(json.dumps(native(structure["pae"]), ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()
                entry["pae_values_sha256"] = structure["pae"].get("values_ref")
            embedded.append(entry)
        provenance["embedded_structures"] = embedded
        provenance["summary"] = payload["summary"]
        payload["provenance"] = provenance
        from ..utils.presentation import portable_metadata
        # Resolve/embed native files first, then sanitize the presentation copy.
        # This never rewrites source CSVs, coordinates, PAE or runtime config.
        payload = portable_metadata(payload, root=self.root)
        all_table, provenance = payload["candidates"], payload["provenance"]
        files = {"report.html": self._offline_document(payload, assets)}
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=all_table["columns"])
        from ..utils.spreadsheet import csv_cell, csv_headers
        writer.writerow(dict(zip(all_table["columns"], csv_headers(all_table["columns"]))))
        writer.writerows([{key: csv_cell(json.dumps(native(value), ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value)
                          for key, value in row.items()} for row in all_table["rows"]])
        files["candidates.csv"] = ("\ufeff" + stream.getvalue()).encode()
        files["sequences.fasta"] = self._fasta(all_table["rows"]).encode()
        retained_note = ("핵심 보관본입니다. 서열·지표·모든 최종 구조를 유지하며, PAE 배열·MSA·중간 파일은 용량 정리로 삭제됐습니다.\n"
                         "이 상세 자료가 필요하면 별도 새 분석에서 전체 계산을 다시 실행하세요. 원래 미평가와 삭제된 자료는 구분합니다.\n") if self.retained else ""
        viewer_features = "표 정렬, MolStar 구조 탐색과 선택 예측 구조 저장" if self.retained else "표 정렬, MolStar 구조 탐색, PAE와 선택 예측 구조 저장"
        files["README.txt"] = ("NanoCoFold 전달용 보고서\n\nZIP을 모두 풀고 report.html을 Chrome, Edge 또는 Firefox에서 여세요.\n"
            f"인터넷·서버·Python·GPU 설치 없이 {viewer_features}을 사용할 수 있습니다.\n" + retained_note +
            "브라우저의 WebGL이 필요합니다. 미평가·실패와 전체 모델을 포함하며 화면 필터로 데이터를 삭제하지 않습니다.\n"
            "MSA, 로그, 가중치, 중간 계산 파일은 포함하지 않습니다. 재실행이나 전체 계산 감사를 위해 전체 원본 ZIP을 별도로 보관하세요.\n"
            "저장된 분석 결과를 공유용으로 구성했습니다. 결합·친화도·KD는 실험으로 확인하세요.\n"
            "delivery_manifest.json은 이 패키지의 파일 및 내장 구조/PAE 해시와 출처를 기록합니다.\n").encode()
        files["README.txt"] += "서버 경로는 상대 경로 또는 [local path]로 표시합니다. 원본 계산 기록은 전체 원본 ZIP에서 확인하세요.\n".encode()
        english = ("NanoCoFold shareable report\n\nExtract the entire ZIP, then open report.html in Chrome, Edge or Firefox.\n"
                   "Choose Language in the report to switch between English and Korean. No Python, server, GPU installation or internet is required.\n"
                   "Check candidate/model counts and exclusions, select a candidate and observation, then inspect its structure and PAE.\n"
                   "3D viewing requires WebGL. Use a current browser to view PAE. Tables and structures remain available if the PAE view is unsupported or too large.\n"
                   "All stored rows, including failed and unevaluated records, remain in the package. Display filters do not remove data.\n"
                   "PAE values retain their full precision and the residue axes for each structure.\n"
                   + ("This is a reduced retention archive: sequences, metrics and final coordinates remain; PAE arrays, MSA and intermediate files were removed by explicit retention.\n" if self.retained else "")
                   + "MSA, logs, weights and intermediate calculation files are omitted. Keep a separate full native archive for reruns and audits.\n"
                   "Confirm binding, affinity and KD experimentally.\n"
                   "delivery_manifest.json lists file hashes, sources, and the records for embedded coordinates and PAE.\n"
                   "Server paths are shown as relative paths or [local path]. Original calculation records remain in the full native archive.\n\n--- 한국어 ---\n\n")
        files["README.txt"] = english.encode() + files["README.txt"]
        for name in ("3DMOL_LICENSE.txt", "MOLSTAR_LICENSE.txt", "PLOTLY_LICENSE.txt", "vendor_manifest.json"):
            content, evidence = _resolve_viewer_asset({}, name, asset_root=ASSET_ROOT, local_only=True)
            if not evidence["ok"]:
                raise OSError(evidence["blocker"])
            files["licenses/" + name] = content
        receipt = {"summary": payload["summary"], "provenance": provenance,
                   "all_models_included": True, "offline": True,
                   "retention": "minimal" if self.retained else "full",
                   "capabilities": payload["summary"].get("capabilities"),
                   "omitted": ["MSA", "logs", "weights", "intermediate_files"] + (["PAE_arrays_removed_by_retention"] if self.retained else [])}
        return package_report(files, native(receipt)), "application/zip", "report_delivery.zip"


def _offline_fallback(encoded):
    return '''<!doctype html><html lang="ko"><meta charset="utf-8"><title>NanoCoFold offline</title>
<style>body{font:16px system-ui;margin:2rem}table{border-collapse:collapse}td,th{padding:.5rem;border:1px solid #ddd}input{padding:.5rem}</style>
<h1>NanoCoFold 결과</h1><p>예측 점수는 결합 확률이나 KD가 아닙니다. 모든 후보와 출처를 포함합니다.</p>
<label>검색 <input id="query"></label><p id="count"></p><div style="overflow:auto"><table id="table"></table></div>
<script id="data" type="application/json">''' + encoded + '''</script><script>
const data=JSON.parse(document.getElementById('data').textContent),table=document.getElementById('table');
function searchText(r){return [r.candidate_id,r.sequence,r.exclusion_reason,...(r.exclusion_reasons||[]),r.error].filter(v=>v!==null&&v!==undefined).join(' ').toLowerCase();}
function render(){const q=document.getElementById('query').value.toLowerCase();const rows=data.candidates.rows.filter(r=>searchText(r).includes(q));
table.replaceChildren();const head=document.createElement('tr');for(const key of data.candidates.columns){const th=document.createElement('th');th.textContent=key;head.append(th)}table.append(head);
for(const row of rows){const tr=document.createElement('tr');for(const key of data.candidates.columns){const td=document.createElement('td');td.textContent=row[key]===null?'미평가':typeof row[key]==='object'?JSON.stringify(row[key]):String(row[key]??'');tr.append(td)}table.append(tr)}document.getElementById('count').textContent=rows.length+' / '+data.candidates.total;}
document.getElementById('query').addEventListener('input',render);render();</script></html>'''
