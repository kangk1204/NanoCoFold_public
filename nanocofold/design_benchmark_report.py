"""Read-only portable reports for paired de novo design campaigns.

Generation records and predictor observations have different grains.
This module preserves both, including failed arms, without launching any tool or
recomputing a stored scientific metric. Coordinates and exact PAE arrays are
content addressed so shared structures are embedded only once.
"""
from __future__ import annotations

import base64
import csv
import gzip
import hashlib
import io
import json
import math
import re
import statistics
from pathlib import Path

from .config import ASSET_ROOT
from .platform.paths import contained, native, sha256
from .platform.results import ResultReader
from .report_delivery import package_report

LEGACY_ENGINES = ("colabfold", "boltz2", "chai1")
AVAILABLE_ENGINES = (*LEGACY_ENGINES, "opendde")
# Compatibility for callers that describe the original campaign protocol.
ENGINES = LEGACY_ENGINES
METRICS = ("ipsae", "iptm", "interface_pae", "interface_plddt", "pdockq2",
           "hotspot_contact_frac", "n_clashes")
SOURCE_FILES = ("run_config.yaml", "run_state.json", "platform_status.json",
                "case_receipt.json", "failure_receipt.json",
                "02_design/all_generated_candidates.csv", "02_design/design_candidates.csv",
                "02_design/engine_info.json", "02_comparison/comparison_status.json",
                "02_design/rfantibody/02_design/all_generated_candidates.csv",
                "02_design/rfantibody/all_generated_candidates.csv",
                "02_design/boltzgen/all_generated_candidates.csv",
                "02_comparison/comparison_identity.json", "02_comparison/comparison_candidates.csv",
                "03_score/scored_candidates.csv", "05_final/selected_candidates.csv",
                "01_msa/api_comparison_acquisition.json", "01_target/residue_mapping.csv")


def _json(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"Expected JSON object: {path.name}")
    return data


def _number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None


def _csv_rows(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    from .platform.results import _cell
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, strict=True)
        keys = reader.fieldnames
        if not keys or any(not key for key in keys) or len(set(keys)) != len(keys):
            raise ValueError(f"Invalid CSV header: {path.name}")
        rows = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Invalid CSV row: {path.name}")
            rows.append({key: _cell(value, key) for key, value in row.items()})
        return rows


def _text_data(rows: list[dict]) -> bytes:
    columns = list(dict.fromkeys(key for row in rows for key in row))
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: json.dumps(native(value), ensure_ascii=False, allow_nan=False)
                         if isinstance(value, (dict, list)) else value for key, value in row.items()})
    return ("\ufeff" + stream.getvalue()).encode()


def _source_id(row):
    identifier = str(row.get("source_sequence_id") or row.get("design_id") or row.get("candidate_id") or "")
    for engine in AVAILABLE_ENGINES:
        if identifier.startswith(engine + "__"):
            return identifier[len(engine) + 2:]
    return identifier


def _declared_engines(inputs):
    protocol = inputs.get("protocol", {})
    if not isinstance(protocol, dict):
        raise ValueError("Campaign protocol must be an object")
    engines = protocol.get("followup_engines", list(LEGACY_ENGINES))
    if (not isinstance(engines, list) or not engines
            or any(not isinstance(engine, str) or engine not in AVAILABLE_ENGINES for engine in engines)
            or len(set(engines)) != len(engines)):
        raise ValueError("Campaign followup_engines must be a unique supported engine list")
    return tuple(engines)


def _case_observations(case, candidates, engines):
    """Resolve a runner's canonical intent ledger without inventing attempts."""
    if "observations" not in case:
        return None
    if not isinstance(case["observations"], list):
        raise ValueError("Canonical observation ledger must be a list")
    states = {"unstarted": "not_started", "not_started": "not_started",
              "running": "running", "interrupted": "interrupted",
              "failed": "failed", "complete": "complete"}
    expected = {(candidate["source_sequence_id"], engine)
                for candidate in candidates.values() if candidate["generation_complete"] for engine in engines}
    observations = {}
    for observation in case["observations"]:
        if not isinstance(observation, dict):
            raise ValueError("Canonical observation must be an object")
        source_id, engine, identifier = (observation.get(key) for key in ("source_sequence_id", "engine", "id"))
        status = observation.get("status")
        if (not all(isinstance(value, str) for value in (source_id, engine, identifier))
                or (source_id, engine) not in expected or identifier != engine + "__" + source_id
                or (source_id, engine) in observations or not isinstance(status, str) or status not in states):
            raise ValueError("Canonical observation identity, coverage or status differs from frozen candidates")
        observations[source_id, engine] = {**observation, "prediction_status": states[status]}
    if set(observations) != expected:
        raise ValueError("Canonical observation coverage differs from frozen candidate/engine requests")
    return observations


def _model_counts(models):
    """Count attempted and unfinished states without classifying interruption as failure."""
    return {"model_records": len(models),
            "models_attempted": sum(row["prediction_status"] in {"complete", "failed", "interrupted", "running"} for row in models),
            "models_evaluated": sum(row["evaluated"] for row in models),
            "models_eligible": sum(row["final_eligible"] is True for row in models),
            "models_failed": sum(row["prediction_status"] == "failed"
                                 or row["prediction_status"] == "complete" and bool(row.get("error")) for row in models),
            "models_interrupted": sum(row["prediction_status"] == "interrupted" for row in models),
            "models_running": sum(row["prediction_status"] == "running" for row in models),
            "models_not_started": sum(row["prediction_status"] == "not_started" for row in models)}


def _campaign_timing(status, *, opendde_addition):
    """Retain clock definitions separately from the original numeric evidence."""
    legacy_elapsed = "runner_entry_through_final_status_checkpoint"
    legacy_timestamps = "post_validation_status_checkpoint_and_final_status_checkpoint"
    elapsed_scope = status.get("elapsed_scope", legacy_elapsed if opendde_addition else None)
    timestamp_scope = status.get("timestamp_scope", legacy_timestamps if opendde_addition else None)
    notes = []
    if opendde_addition and ("elapsed_scope" not in status or "timestamp_scope" not in status):
        notes.append("기존 추가 실행기 정의에 따른 측정 범위입니다.")
    if elapsed_scope == legacy_elapsed:
        notes.append("경과시간은 실행기 진입부터 마지막 상태 기록까지이며 입력 검증·확인·자원 대기·예측·점수 계산을 포함하고 이후 산출물 봉인은 제외합니다.")
    if timestamp_scope == legacy_timestamps:
        notes.append("시작 UTC는 입력 검증 뒤 상태 기록 시점이고 종료 UTC는 최종 상태 기록 시점입니다.")
    elif timestamp_scope == "runner_entry_and_final_status_checkpoint":
        notes.append("시작·종료 UTC는 실행기 진입과 최종 상태 기록 시점입니다.")
    if elapsed_scope is not None or timestamp_scope is not None:
        notes.append("원래 값은 보존하며 UTC 차이로 경과시간을 재계산하지 않습니다.")
    else:
        notes.append("측정 범위는 원본에 미기록입니다. 원래 경과시간·UTC 값은 보존하며 UTC 차이로 재계산하지 않습니다.")
    return {"elapsed_scope": elapsed_scope, "timestamp_scope": timestamp_scope, "timing_note": " ".join(notes)}


def _origin(row):
    text = str(row.get("parent_engine") or row.get("design_origin") or row.get("source") or _source_id(row)).lower()
    return "rfantibody" if "rfantibody" in text else "boltzgen" if "boltzgen" in text else text or "unrecorded"


def _candidate_records(run: Path, case_id: str) -> list[dict]:
    path = run / "02_design/all_generated_candidates.csv"
    records = _csv_rows(path)
    # A later generator can fail while dispatch's combined table contains only
    # the earlier engine. Supplement its native failure records, not just when
    # the combined table is completely absent.
    known = {_source_id(row): row for row in records}
    for engine, relative in (("rfantibody", "02_design/rfantibody/02_design/all_generated_candidates.csv"),
                                 ("rfantibody", "02_design/rfantibody/all_generated_candidates.csv"),
                                 ("boltzgen", "02_design/boltzgen/all_generated_candidates.csv")):
        for row in _csv_rows(run / relative):
            identifier = str(row.get("design_id") or "")
            if not identifier:
                raise ValueError("Generated record has no source ID")
            row = {**row, "parent_engine": engine,
                   "design_id": identifier if identifier.startswith(engine + "__") else engine + "__" + identifier}
            identifier = _source_id(row)
            if identifier in known:
                if known[identifier].get("sequence") != row.get("sequence"):
                    raise ValueError(f"Combined and native generated sequences differ: {case_id}")
                continue
            records.append(row)
            known[identifier] = row
    result, seen, counts = [], set(), {"rfantibody": 0, "boltzgen": 0}
    for row in records:
        identifier = _source_id(row)
        if not identifier or identifier in seen:
            raise ValueError(f"Duplicate/missing generated source identity: {case_id}")
        seen.add(identifier)
        origin = _origin(row)
        counts[origin] = counts.get(origin, 0) + 1
        alias = {"rfantibody": "RF", "boltzgen": "BG"}.get(origin, "VHH") + f"-{counts[origin]:02d}"
        sequence = row.get("sequence")
        valid = isinstance(sequence, str) and bool(re.fullmatch("[ACDEFGHIKLMNPQRSTVWY]+", sequence))
        error = row.get("generation_error")
        complete = valid and not error and row.get("generation_status") not in {"failed", "interrupted"}
        result.append(native({**row, "case_id": case_id, "source_sequence_id": identifier,
                              "candidate_record_id": case_id + "__" + identifier,
                              "candidate_label": alias, "design_origin": origin,
                              "sequence_sha256": hashlib.sha256(sequence.encode()).hexdigest() if valid else None,
                              "sequence_length": len(sequence) if valid else None,
                              "generation_complete": complete}))
    return result


def collect_design_benchmark(source: Path) -> dict:
    """Extract canonical ledgers and exact viewer data, with no inference."""
    source = Path(source).resolve()
    campaign_digests = {name: sha256(contained(source, name, exists=True))
                        for name in ("benchmark_status.json", "inputs_manifest.json")}
    for name in ("execution_identity.json", "augmentation_identity.json"):
        if (source / name).is_file():
            campaign_digests[name] = sha256(contained(source, name, exists=True))
    status = _json(contained(source, "benchmark_status.json", exists=True))
    inputs = _json(contained(source, "inputs_manifest.json", exists=True))
    engines = _declared_engines(inputs)
    if status.get("schema_version") != 1 or not isinstance(status.get("cases"), list):
        raise ValueError("Unsupported benchmark status schema")
    if not isinstance(inputs.get("cases"), list):
        raise ValueError("Missing target inputs")
    targets, target_ids = [], set()
    for item in inputs["cases"]:
        identifier = item.get("pdb_id")
        target = item.get("target", {})
        sequence = target.get("sequence")
        if not isinstance(identifier, str) or not re.fullmatch("[A-Za-z0-9_-]+", identifier) or identifier in target_ids or not isinstance(sequence, str) or not sequence:
            raise ValueError("Missing/duplicate target identity or sequence")
        digest = hashlib.sha256(sequence.encode()).hexdigest()
        if target.get("sequence_sha256") != digest:
            raise ValueError(f"Target sequence hash changed: {identifier}")
        target_ids.add(identifier)
        targets.append({"pdb_id": identifier, "name": item.get("name"), "sequence": sequence,
                        "sequence_sha256": digest, "sequence_length": len(sequence),
                        "hotspots_seq_positions": item.get("hotspots_seq_positions", []),
                        "source_url": item.get("source_url")})
    target_index = {item["pdb_id"]: item for item in targets}
    cases, generated, models, sources, issues = [], [], [], [], []
    coordinates, pae_arrays, seen_cases, seen_models = {}, {}, set(), set()
    design_coordinates = {}
    pose_configs = {}
    for case in status["cases"]:
        pdb_id, arm = case.get("pdb_id"), case.get("arm")
        if pdb_id not in target_ids or arm not in {"off", "on"}:
            raise ValueError("Case is outside the exact paired target protocol")
        case_id = pdb_id + "__" + arm
        if case_id in seen_cases:
            raise ValueError("Duplicate target/arm case")
        seen_cases.add(case_id)
        relative = case.get("run_dir")
        if not isinstance(relative, str) or not relative or Path(relative).is_absolute():
            raise ValueError("Case run_dir must be a contained relative path")
        run = contained(source, relative)
        candidate_rows = _candidate_records(run, case_id) if run.is_dir() else []
        for row in candidate_rows:
            row.update(pdb_id=pdb_id, arm=arm)
        generated.extend(candidate_rows)
        candidate_index = {row["source_sequence_id"]: row for row in candidate_rows}
        canonical = _case_observations(case, candidate_index, engines)
        run_issues, rows = [], []
        reader = ResultReader(run) if run.is_dir() and (run / "run_config.yaml").is_file() else None
        if reader is not None:
            from .pose_assessment import settings as pose_settings, DEFAULTS as pose_defaults
            frozen_policy = pose_settings(reader.config)
            pose_configs[case_id] = {"rank": {"pose_assessment": {key: frozen_policy[key] for key in pose_defaults}}}
            for candidate in candidate_rows:
                value = candidate.get("design_structure_path") or candidate.get("cif_path")
                if not isinstance(value, str) or not value:
                    continue
                try:
                    path = reader.artifact_path(value)
                    content = path.read_text(encoding="utf-8")
                    digest = sha256(path)
                    if hashlib.sha256(content.encode()).hexdigest() != digest:
                        raise ValueError("Generation coordinate bytes changed")
                    design_coordinates[candidate["candidate_record_id"]] = {
                        "content": content, "format": "pdb" if path.suffix.lower() == ".pdb" else "cif",
                        "sha256": digest, "target_chain": candidate.get("design_target_chain") or candidate.get("target_chain"),
                        "binder_chain": candidate.get("design_binder_chain") or candidate.get("binder_chain"),
                        "reference_stage": candidate.get("design_reference_stage") or candidate.get("structure_stage")}
                    sources.append({"case_id": case_id, "candidate_record_id": candidate["candidate_record_id"],
                                    "kind": "generation_stage_coordinate", "path": str(path.relative_to(source)),
                                    "sha256": digest, "size": path.stat().st_size})
                except (OSError, ValueError, KeyError) as exc:
                    # The independent prediction remains valid when its design
                    # reference is unavailable. Disclose this diagnostic gap.
                    run_issues.append(f"{candidate['candidate_record_id']}: generation comparison unavailable: {exc}")
            if (run / "03_score/scored_candidates.csv").is_file():
                normalized = {row["candidate_id"]: row for row in reader.all_rows()}
                # Preserve original scalar scientific fields. Legacy mixed
                # success/failure CSVs may omit final_evaluated on successful
                # rows; the reader's display defaults must not erase evidence.
                for index, raw in enumerate(_csv_rows(run / "03_score/scored_candidates.csv")):
                    identifier = reader._identifier(raw, index)
                    rows.append({**normalized.get(identifier, {}), **raw, "candidate_id": identifier})
            if not rows:
                rows = _csv_rows(run / "02_comparison/comparison_candidates.csv")
            run_issues.extend(reader.issues)
        if canonical is not None:
            for row in rows:
                identifier = str(row.get("candidate_id") or row.get("design_id") or "")
                observation = canonical.get((_source_id(row), row.get("predictor_engine") or row.get("engine")))
                prediction_state = row.get("prediction_status") or ("failed" if row.get("prediction_error") else "complete")
                if observation is None or identifier != observation["id"] or prediction_state != observation["prediction_status"]:
                    raise ValueError("Canonical observation and stored CSV identity/status differ")
        # An MSA or generator failure may precede all predictor submissions.
        # These explicit records describe unstarted requests, not completed models.
        for candidate in candidate_rows:
            if not candidate["generation_complete"]:
                continue
            observed = {(_source_id(row), row.get("predictor_engine") or row.get("engine")) for row in rows}
            for engine in engines:
                if (candidate["source_sequence_id"], engine) not in observed:
                    observation = canonical.get((candidate["source_sequence_id"], engine)) if canonical is not None else None
                    prediction_state = observation["prediction_status"] if observation else "not_started"
                    error = observation.get("error") if observation else None
                    if prediction_state == "not_started":
                        error = error or case.get("error") or "No stored prediction observation"
                    rows.append({"candidate_id": engine + "__" + candidate["source_sequence_id"],
                                 "source_sequence_id": candidate["source_sequence_id"], "sequence": candidate["sequence"],
                                 "predictor_engine": engine, "prediction_status": prediction_state, "final_evaluated": False,
                                 "error": error, "model_record_kind": "planned_unstarted" if prediction_state == "not_started" else "receipt_observation"})
        for index, raw in enumerate(rows):
            row = native(raw)
            identifier = str(row.get("candidate_id") or row.get("design_id") or "")
            engine = row.get("predictor_engine") or row.get("engine")
            if not identifier or engine not in engines:
                raise ValueError(f"Model observation has missing identity/engine: {case_id}")
            from .utils.seq import validate_candidate_id
            validate_candidate_id(identifier)
            model_id = case_id + "__" + identifier
            if model_id in seen_models:
                raise ValueError("Duplicate global model identity")
            seen_models.add(model_id)
            source_id = _source_id(row)
            candidate = candidate_index.get(source_id)
            if candidate is not None and row.get("sequence") != candidate.get("sequence"):
                raise ValueError(f"Generated and evaluated exact sequences differ: {model_id}")
            prediction_status = row.get("prediction_status") or ("failed" if row.get("prediction_error") else "complete")
            error = row.get("prediction_error") or row.get("error")
            recorded = row.get("final_evaluated")
            evaluated = prediction_status == "complete" and not error and recorded is not False and any(_number(row.get(k)) is not None for k in METRICS)
            # Missing metrics stay null. Never turn a missing predictor into 0.
            row.update(id=model_id, case_id=case_id, pdb_id=pdb_id, arm=arm,
                       candidate_id=identifier, source_sequence_id=source_id,
                       candidate_label=candidate["candidate_label"] if candidate else source_id,
                       design_origin=candidate["design_origin"] if candidate else _origin(row),
                       predictor_engine=engine, evaluated=evaluated,
                       recorded_final_evaluated=recorded, final_evaluated=evaluated,
                       final_eligible=row.get("final_eligible") if evaluated and isinstance(row.get("final_eligible"), bool) else None,
                       prediction_status=prediction_status, error=error,
                       model_record_kind=row.get("model_record_kind", "stored_observation"),
                       structure_key=None, pae_key=None,
                       target_chain=str(row.get("target_chain") or "A"), binder_chain=str(row.get("binder_chain") or "B"))
            if not evaluated:
                for key in METRICS:
                    row[key] = None
            else:
                for key in METRICS:
                    row[key] = _number(row.get(key))
            from .utils.metric_provenance import metric_provenance
            row["metric_provenance"] = metric_provenance(row)
            if reader is not None and prediction_status == "complete" and row.get("cif_path"):
                try:
                    path = reader.artifact_path(row["cif_path"])
                    content = path.read_bytes()
                    text = content.decode("utf-8")
                    if text.encode("utf-8") != content:
                        raise ValueError("Coordinate UTF-8 bytes changed")
                    digest = hashlib.sha256(content).hexdigest()
                    coordinates[digest] = {"content": text, "format": "pdb" if path.suffix.lower() == ".pdb" else "cif"}
                    row.update(structure_key=digest, structure_sha256=digest)
                    sources.append({"case_id": case_id, "model_id": model_id, "kind": "coordinate",
                                    "path": str(path.relative_to(source)), "sha256": digest, "size": len(content)})
                    _attach_pae(reader, path, row, pae_arrays, sources, source)
                except (OSError, ValueError, KeyError) as exc:
                    run_issues.append(f"{model_id}: structure/PAE unavailable: {exc}")
            models.append(row)
        for name in SOURCE_FILES:
            path = contained(run, name)
            if path.is_file():
                sources.append({"case_id": case_id, "kind": "source_receipt", "path": str(path.relative_to(source)),
                                "sha256": sha256(path), "size": path.stat().st_size})
        api = _json(run / "01_msa/api_comparison_acquisition.json") if (run / "01_msa/api_comparison_acquisition.json").is_file() else {}
        api_times = [_number(batch.get("elapsed_seconds")) for batch in api.get("batches", {}).values()]
        api_times = [value for value in api_times if value is not None]
        engine_info = _json(run / "02_design/engine_info.json") if (run / "02_design/engine_info.json").is_file() else {}
        config = reader.config if reader is not None else {}
        design = config.get("design", {})
        screen = config.get("screen", {})
        conditions = {"run_seed": config.get("run", {}).get("seed"), "mock": config.get("run", {}).get("mock"),
                      "design": {key: design.get(key) for key in ("engine", "hotspot_mode", "num_designs", "refinement", "filters", "mpnn", "rf2")},
                      "comparison": screen.get("comparison"), "restraint": screen.get("restraint"),
                      "msa": {key: screen.get("colabfold", {}).get(key) for key in ("msa_source", "msa_mode", "msa_max_sequences")}}
        case_models = [row for row in models if row["case_id"] == case_id]
        if case.get("status") == "complete" and any(row["prediction_status"] == "not_started" for row in case_models):
            run_issues.append(f"{case_id}: marked complete with missing predictor observations")
        cases.append(native({**case, "case_id": case_id, "target_name": target_index[pdb_id]["name"],
                             "generated_records": len(candidate_rows),
                             "generated_complete": sum(row["generation_complete"] for row in candidate_rows),
                             "unique_sequences": len({row["sequence"] for row in candidate_rows if row["generation_complete"]}),
                             **_model_counts(case_models),
                             "network_requests_this_run": api.get("recorded_network_requests_this_run"),
                             "source_ticket_lifetime_requests": api.get("public_api_requests"),
                             "conditions": conditions, "generation_engine_info": engine_info,
                             "api_msa_elapsed_seconds": sum(api_times) if api_times else None,
                             "issues": list(dict.fromkeys(run_issues))}))
        issues.extend(run_issues)
    models.sort(key=lambda row: (row.get("ipsae") is None, -row["ipsae"] if row.get("ipsae") is not None else 0, row["id"]))
    for name in campaign_digests:
        path = source / name
        sources.append({"kind": "campaign_receipt", "path": name, "sha256": campaign_digests[name], "size": path.stat().st_size})
    if any(sha256(contained(source, receipt["path"], exists=True)) != receipt["sha256"] for receipt in sources):
        raise ValueError("Source evidence changed during report extraction; use a stable preserved snapshot")
    timing_groups = {}
    for case in cases:
        observed = {"case_seconds": case.get("elapsed_seconds"), **(case.get("timings") or {})}
        for stage, value in observed.items():
            if _number(value) is not None and value >= 0:
                timing_groups.setdefault((case["arm"], stage), []).append(value)
    timing_summary = [{"arm": arm, "stage": stage, "observations": len(values),
                       "total_seconds": sum(values), "mean_seconds": statistics.mean(values),
                       "median_seconds": statistics.median(values), "min_seconds": min(values),
                       "max_seconds": max(values)} for (arm, stage), values in sorted(timing_groups.items())]
    summary = {"status": status.get("status"), "expected_cases": status.get("expected_cases"),
               "expected_candidates": status.get("expected_candidates"), "expected_models": status.get("expected_models"),
               "recorded_cases": len(cases), "terminal_cases": sum(c.get("status") in {"complete", "failed", "partial_failure", "interrupted", "cancelled"} for c in cases),
               "generated_records": len(generated), "generated_complete": sum(row["generation_complete"] for row in generated),
               "unique_sequences": len({row["sequence"] for row in generated if row["generation_complete"]}),
               **_model_counts(models),
               **_campaign_timing(status, opendde_addition="augmentation_identity.json" in campaign_digests and engines == ("opendde",)),
               "started_utc": status.get("started_utc"), "completed_utc": status.get("completed_utc"),
               "elapsed_seconds": _number(status.get("elapsed_seconds")), "unique_coordinate_files": len(coordinates),
               "unique_pae_arrays": len(pae_arrays), "issues": list(dict.fromkeys(issues))}
    return native({"schema_version": 1, "kind": "de_novo_hotspot_paired", "summary": summary,
                   "engines": list(engines),
                   "targets": targets, "cases": cases, "timing_summary": timing_summary,
                   "campaign_protocol": {key: value for key, value in inputs.items() if key != "cases"},
                   "generated_candidates": generated, "models": models,
                   "coordinates": coordinates, "design_coordinates": design_coordinates, "pose_configs": pose_configs,
                   "pae_arrays": pae_arrays, "source_artifacts": sources,
                   "interpretation": {"score_aggregation": "none", "default_sort": "ipsae_desc_missing_last",
                                      "native_DockQ": "not_applicable_to_novel_candidate_sequences",
                                      "binding_validation": "not_established", "new_inference_from_report": 0,
                                      "new_network_requests_from_report": 0}})


def collect_augmented_design_benchmark(source: Path, augmentation: Path) -> dict:
    """Join preserved three-engine results with an OpenDDE-only addition.

    The imported candidate ledger describes the same historical generation.
    Model dictionaries and content-addressed viewer assets retain their exact
    stored values; each computation's receipts and timings keep their scope.
    """
    roots = {"original_3engine": Path(source).resolve(), "opendde_addition": Path(augmentation).resolve()}
    if roots["original_3engine"] == roots["opendde_addition"]:
        raise ValueError("Original campaign and augmentation must have separate source roots")
    original = collect_design_benchmark(roots["original_3engine"])
    addition = collect_design_benchmark(roots["opendde_addition"])
    if tuple(original["engines"]) != LEGACY_ENGINES or addition["engines"] != ["opendde"]:
        raise ValueError("Augmentation requires the legacy three engines and an OpenDDE-only addition")
    if original["targets"] != addition["targets"]:
        raise ValueError("Augmentation exact target inputs differ from the original campaign")
    if {row["case_id"] for row in original["cases"]} != {row["case_id"] for row in addition["cases"]}:
        raise ValueError("Augmentation target/arm case coverage differs from the original campaign")

    def candidate_identity(payload):
        result = {}
        for row in payload["generated_candidates"]:
            key = row["candidate_record_id"]
            if key in result or not row["generation_complete"] or not row["sequence_sha256"]:
                raise ValueError("Augmentation requires unique, complete frozen source candidates")
            result[key] = (row["sequence"], row["sequence_sha256"], row["case_id"], row["source_sequence_id"])
        if not result:
            raise ValueError("Augmentation requires frozen source candidates")
        return result

    candidates = candidate_identity(original)
    if candidates != candidate_identity(addition):
        raise ValueError("Augmentation exact candidate record IDs, sequences or SHA256 differ")
    for payload in (original, addition):
        expected = {(row["case_id"], row["source_sequence_id"], engine)
                    for row in payload["generated_candidates"] for engine in payload["engines"]}
        observed = [(row["case_id"], row["source_sequence_id"], row["predictor_engine"]) for row in payload["models"]]
        if len(observed) != len(set(observed)) or set(observed) != expected:
            raise ValueError("Augmentation model observations must match each frozen candidate/engine once")
        if type(payload["summary"]["expected_models"]) is not int or payload["summary"]["expected_models"] != len(expected):
            raise ValueError("Augmentation expected model count differs from its frozen candidate/engine coverage")
    model_ids = [row["id"] for payload in (original, addition) for row in payload["models"]]
    if len(model_ids) != len(set(model_ids)):
        raise ValueError("Augmentation has duplicate global model identities")

    # Preserve the old asset dictionaries, including their keys and labels.
    coordinates, pae_arrays = dict(original["coordinates"]), dict(original["pae_arrays"])
    for key, value in addition["coordinates"].items():
        if key in coordinates and coordinates[key] != value:
            raise ValueError("Augmentation coordinate content-address collision")
        coordinates.setdefault(key, value)
    for key, value in addition["pae_arrays"].items():
        if key in pae_arrays and any(pae_arrays[key][field] != value[field]
                                     for field in ("encoding", "dimension", "unit", "sha256", "bytes", "data")):
            raise ValueError("Augmentation PAE content-address collision")
        pae_arrays.setdefault(key, value)

    models = original["models"] + addition["models"]
    models.sort(key=lambda row: (row.get("ipsae") is None, -row["ipsae"] if row.get("ipsae") is not None else 0, row["id"]))
    sources, cases, timing_summary, timing_scopes, source_campaigns = [], [], [], [], {}
    for scope, payload in (("original_3engine", original), ("opendde_addition", addition)):
        root = str(roots[scope])
        sources.extend({**row, "source_scope": scope, "source_root": root} for row in payload["source_artifacts"])
        cases.extend({**row, "source_scope": scope, "source_root": root} for row in payload["cases"])
        timing_summary.extend({**row, "source_scope": scope} for row in payload["timing_summary"])
        timing_scopes.append({"source_scope": scope, "source_root": root,
                              "engines": payload["engines"], "status": payload["summary"]["status"],
                              **{key: payload["summary"][key] for key in ("started_utc", "completed_utc", "elapsed_seconds",
                                                                       "elapsed_scope", "timestamp_scope", "timing_note")}})
        source_campaigns[scope] = {"source_root": root, "kind": payload["kind"], "engines": payload["engines"],
                                   "summary": payload["summary"], "campaign_protocol": payload["campaign_protocol"]}
    for receipt in sources:
        if sha256(contained(Path(receipt["source_root"]), receipt["path"], exists=True)) != receipt["sha256"]:
            raise ValueError("Source evidence changed during augmentation extraction")
    issues = list(dict.fromkeys(original["summary"]["issues"] + addition["summary"]["issues"]))
    states = [payload["summary"]["status"] for payload in (original, addition)] + [row["status"] for row in cases]
    terminal = {"complete", "failed", "partial_failure", "interrupted", "cancelled"}
    complete = (all(state == "complete" for state in states) and not issues
                and all(row["evaluated"] for row in models))
    status = "complete" if complete else "partial_failure" if all(state in terminal for state in states) else "incomplete"
    summary = {**original["summary"], "status": status,
               "expected_models": sum(payload["summary"]["expected_models"] for payload in (original, addition)),
               "recorded_cases": len({row["case_id"] for row in cases}), "source_case_records": len(cases),
               "terminal_cases": sum(all(row["status"] in terminal for row in cases if row["case_id"] == case_id)
                                     for case_id in {row["case_id"] for row in cases}),
               **_model_counts(models),
               "legacy_models_preserved": len(original["models"]), "added_model_records": len(addition["models"]),
               "legacy_models_evaluated": original["summary"]["models_evaluated"],
               "added_models_evaluated": addition["summary"]["models_evaluated"],
               "new_candidates_generated": 0, "elapsed_seconds": None, "started_utc": None, "completed_utc": None,
               "elapsed_scope": "separate_source_campaigns", "timestamp_scope": "separate_source_campaigns",
               "timing_note": "각 계산의 경과시간·UTC 측정 범위와 값을 각각 보존하며 합산하거나 UTC 차이로 재계산하지 않습니다.",
               "unique_coordinate_files": len(coordinates),
               "unique_pae_arrays": len(pae_arrays), "issues": issues}
    return {**original, "kind": "de_novo_hotspot_paired_opendde_augmented", "engines": list(AVAILABLE_ENGINES),
            "summary": summary, "models": models, "cases": cases, "coordinates": coordinates, "pae_arrays": pae_arrays,
            "source_artifacts": sources, "source_campaigns": source_campaigns, "timing_scopes": timing_scopes,
            "timing_summary": timing_summary, "augmentation_protocol": addition["campaign_protocol"],
            "interpretation": {**original["interpretation"], "original_models_preserved": True,
                               "new_candidate_generation": 0, "new_msa_acquisition_from_report": 0,
                               "timing_scope": "original_3engine_and_opendde_addition_separate",
                               "binding_validation": "not_established", "training_independence": "not_established",
                               "hotspot_causality": "not_established"}}


def _attach_pae(reader, path, row, arrays, sources, source):
    from . import score_complex
    from .utils import pdb_io
    prediction = reader._prediction_for_path(path)
    associated_sources = []
    for kind, associated in (("pae_source", prediction.scores_npz), ("pae_source", prediction.pae_npz),
                             ("confidence_source", prediction.confidence_json)):
        if associated and associated.is_file():
            associated = reader.artifact_path(associated)
            associated_sources.append({"case_id": row["case_id"], "model_id": row["id"], "kind": kind,
                                       "path": str(associated.relative_to(source)), "sha256": sha256(associated), "size": associated.stat().st_size})
    sources.extend(associated_sources)
    matrix = score_complex.load_pae(prediction)
    if matrix is None:
        return
    labels, _ = score_complex.build_index_map(pdb_io.load_structure(path))
    score_complex.validate_pae(matrix, len(labels))
    binary = matrix.astype("<f8", copy=False).tobytes(order="C")
    digest = hashlib.sha256(binary).hexdigest()
    # Lossless binary compression, not scientific rounding or subsampling.
    arrays.setdefault(digest, {"encoding": "gzip-base64-float64le", "dimension": len(labels), "labels": labels,
                      "unit": "Å", "sha256": digest, "bytes": len(binary),
                      "data": base64.b64encode(gzip.compress(binary, compresslevel=9, mtime=0)).decode()})
    row["pae_key"] = digest
    row["pae_labels"] = labels


def _excel_text_chunks(text):
    """Split receipt text without loss, counting Excel's UTF-16 cell units."""
    chunks, start, units = [], 0, 0
    for index, char in enumerate(text):
        width = 2 if ord(char) > 0xFFFF else 1
        if units + width > 30000:
            chunks.append(text[start:index])
            start, units = index, 0
        units += width
    chunks.append(text[start:])
    return chunks


def _workbook(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    long_fields = []
    for name, rows in (("Models", payload["models"]), ("Generation", payload["generated_candidates"]),
                       ("Cases and timing", payload["cases"]), ("Targets", payload["targets"]),
                       ("Timing summary", payload["timing_summary"]),
                       ("Timing scopes", payload.get("timing_scopes", [])),
                       ("Source campaigns", [{"source_scope": scope, **row} for scope, row in payload.get("source_campaigns", {}).items()]),
                       ("Sources", payload["source_artifacts"]),
                       ("PDB references", payload.get("reference_structures", [])),
                       ("Hotspot assessment", [{"model_id": key, **value} for key, value in payload.get("hotspot_assessments", {}).items()]),
                       ("Candidate review", payload.get("pose_assessment", {}).get("candidates", [])),
                       ("Pose comparisons", payload.get("pose_assessment", {}).get("pairs", [])),
                       ("Reference binding sites", [{"model_id": key, **value} for key, value in payload.get("reference_binding_sites", {}).items()]),
                       ("Visual target alignment", [{"model_id": key, **value} for key, value in payload.get("reference_overlays", {}).items()])):
        sheet = book.create_sheet(name)
        columns = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(columns)
        for row in rows:
            values = [json.dumps(row.get(k), ensure_ascii=False, allow_nan=False) if isinstance(row.get(k), (dict, list)) else row.get(k) for k in columns]
            # Full native artifact manifests exceed Excel's per-cell limit.
            # Preserve them losslessly in a linked chunk ledger, never truncate.
            for index, value in enumerate(values):
                if isinstance(value, str) and len(value.encode("utf-16-le")) // 2 > 32767:
                    digest = hashlib.sha256(value.encode()).hexdigest()
                    chunks = _excel_text_chunks(value)
                    first = len(long_fields) + 2
                    long_fields.extend((name, sheet.max_row + 1, columns[index], digest,
                                        i + 1, len(chunks), chunk) for i, chunk in enumerate(chunks))
                    values[index] = f"Long fields rows {first}-{len(long_fields) + 1}; SHA256 {digest}"
            sheet.append(values)
            for cell, value in zip(sheet[sheet.max_row], values):
                if isinstance(value, str):
                    cell.data_type = "s"
                if isinstance(value, float):
                    cell.number_format = "0.000000"
    if long_fields:
        sheet = book.create_sheet("Long fields")
        sheet.append(("sheet", "row", "field", "sha256", "chunk", "chunks", "text"))
        for values in long_fields:
            sheet.append(values)
            for cell, value in zip(sheet[sheet.max_row], values):
                if isinstance(value, str):
                    cell.data_type = "s"
    for sheet in book:
        sheet.freeze_panes = "A2"
        if sheet.max_column and sheet.cell(1, 1).value is not None:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244B55")
            cell.alignment = Alignment(wrap_text=True)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def build_design_benchmark(source: Path, output: Path, *, reference_manifest: Path | None = None) -> dict:
    """Write a new presentation revision outside immutable campaign evidence."""
    source, output = Path(source).resolve(), Path(output).resolve()
    if output.exists() or output.is_relative_to(source):
        raise ValueError("Choose a new report directory outside the source campaign")
    payload = collect_design_benchmark(source)
    from .design_reference_overlay import attach_reference_overlays
    attach_reference_overlays(payload, reference_manifest)
    return _write_design_presentation(payload, output)


def build_augmented_design_benchmark(source: Path, augmentation: Path, output: Path,
                                     reference_manifest: Path | None = None) -> dict:
    """Package a new presentation without changing either source computation."""
    source, augmentation, output = Path(source).resolve(), Path(augmentation).resolve(), Path(output).resolve()
    if output.exists() or output.is_relative_to(source) or output.is_relative_to(augmentation):
        raise ValueError("Choose a new report directory outside both source campaigns")
    payload = collect_augmented_design_benchmark(source, augmentation)
    from .design_reference_overlay import attach_reference_overlays
    attach_reference_overlays(payload, reference_manifest)
    return _write_design_presentation(payload, output)


def _write_design_presentation(payload: dict, output: Path) -> dict:
    """Render an already verified payload to a new directory without inference."""
    output = Path(output).resolve()
    if output.exists():
        raise ValueError("Choose a new report directory")
    from .pose_integration import assess_benchmark_payload
    payload = assess_benchmark_payload(payload)
    from .binding_site_assessment import hotspot_assessment
    payload["hotspot_assessments"] = {row["id"]: hotspot_assessment(row) for row in payload["models"]}
    from .utils.presentation import portable_metadata
    payload = portable_metadata(payload)
    from .report import _resolve_viewer_asset
    # Use the same integrity/fallback gate as doctor. A reused workspace can
    # have an older notice/manifest even when its viewer bytes are unchanged.
    names = ("molstar.js", "molstar.css", "plotly.min.js",
             "MOLSTAR_LICENSE.txt", "PLOTLY_LICENSE.txt", "vendor_manifest.json")
    assets = {}
    for name in names:
        content, evidence = _resolve_viewer_asset({}, name, asset_root=ASSET_ROOT, local_only=True)
        if not evidence["ok"] or content is None:
            raise OSError(evidence["blocker"])
        assets[name] = content
    encoded = json.dumps(payload, ensure_ascii=False, allow_nan=False).replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    document = _DOCUMENT
    for token, name in (("__MOLSTAR_JS__", "molstar.js"), ("__MOLSTAR_CSS__", "molstar.css"), ("__PLOTLY_JS__", "plotly.min.js")):
        text = assets[name].decode("utf-8")
        document = document.replace(token, re.sub(r"</(script|style)", r"<\\/\1", text, flags=re.I))
    display_script = (Path(__file__).with_name("templates") / "benchmark-viewer-display.js").read_text()
    document = document.replace("__BENCHMARK_VIEWER_DISPLAY_JS__", display_script)
    # Inject user/source data last so literal vendor tokens remain data.
    document = document.replace("__BENCHMARK_DATA__", encoded)
    # Canonical JSON keeps metrics/ledgers, while viewer bytes occur only in HTML.
    ledger = {key: value for key, value in payload.items() if key not in {"coordinates", "design_coordinates", "pae_arrays", "reference_coordinates"}}
    ledger["embedded_design_coordinates"] = [{"candidate_record_id": key, "sha256": value["sha256"],
        "reference_stage": value.get("reference_stage"), "format": value["format"],
        "bytes": len(value["content"].encode())} for key, value in payload["design_coordinates"].items()]
    ledger["embedded_coordinates"] = [{"sha256": key, "format": value["format"], "bytes": len(value["content"].encode())} for key, value in payload["coordinates"].items()]
    ledger["embedded_pae_arrays"] = [{"sha256": key, "dimension": value["dimension"], "bytes": value["bytes"], "encoding": value["encoding"]} for key, value in payload["pae_arrays"].items()]
    ledger["embedded_reference_coordinates"] = [{"sha256": key, "format": value["format"], "bytes": len(value["content"].encode())} for key, value in payload["reference_coordinates"].items()]
    files = {"report.html": document.encode(), "benchmark.json": json.dumps(ledger, ensure_ascii=False, indent=2, allow_nan=False).encode(),
             "models.csv": _text_data(payload["models"]), "generated_candidates.csv": _text_data(payload["generated_candidates"]),
             "cases.csv": _text_data(payload["cases"]), "benchmark.xlsx": _workbook(payload),
             "target_sequences.fasta": "".join(f">{row['pdb_id']} target_sha256={row['sequence_sha256']}\n{row['sequence']}\n" for row in payload["targets"]).encode(),
             "candidate_sequences.fasta": "".join(f">{row['candidate_record_id']} label={row['candidate_label']} origin={row['design_origin']} sha256={row['sequence_sha256']}\n{row['sequence']}\n" for row in payload["generated_candidates"] if row.get("sequence_sha256")).encode(),
             "README.txt": _README.encode()}
    files["reference_alignment.csv"] = _text_data([{"model_id": key, **value} for key, value in payload["reference_overlays"].items()])
    files["hotspot_assessments.csv"] = _text_data([{"model_id": key, **value} for key, value in payload["hotspot_assessments"].items()])
    assessment = payload["pose_assessment"]
    files["pose_assessment.json"] = json.dumps(assessment, ensure_ascii=False, indent=2, allow_nan=False).encode()
    files["candidate_review.csv"] = _text_data(assessment["candidates"])
    files["pose_pairs.csv"] = _text_data(assessment["pairs"])
    files["reference_binding_sites.csv"] = _text_data([{"model_id": key, **value} for key, value in payload.get("reference_binding_sites", {}).items()])
    if payload.get("timing_scopes"):
        files["timing_scopes.csv"] = _text_data(payload["timing_scopes"])
    for name in ("MOLSTAR_LICENSE.txt", "PLOTLY_LICENSE.txt", "vendor_manifest.json"):
        files["licenses/" + name] = assets[name]
    receipt = {"summary": payload["summary"], "source_artifacts": payload["source_artifacts"],
               "interpretation": payload["interpretation"], "offline": True,
               "coordinates_sha256": list(payload["coordinates"]), "pae_binary_sha256": list(payload["pae_arrays"]),
               "reference_coordinates_sha256": list(payload["reference_coordinates"]),
               "reference_sources": payload["reference_sources"],
               "source_campaigns": payload.get("source_campaigns", {}),
               "timing_scopes": payload.get("timing_scopes", []),
               "omitted": ["MSA", "weights", "logs", "intermediate_tensors", "duplicate_case_HTML"],
               "scientific_metrics_recomputed": False, "source_readonly": True}
    receipt["new_presentation_diagnostics"] = ["recorded_hotspot_assessment", "target_heavy_atom_contact_footprint_5A",
        "same_sequence_target_aligned_pose_consensus", "generation_stage_independent_prediction_consistency",
        "within_case_candidate_review_rank"]
    receipt["pose_policy"] = assessment["policy"]
    receipt["new_geometry_diagnostics_computed"] = True
    receipt["recorded_prediction_scores_recomputed"] = False
    archive = package_report(files, receipt)
    output.mkdir(parents=True)
    for name, content in files.items():
        destination = output / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
    archive_path = output / "NanoCoFold_design_hotspot_report.zip"
    archive_path.write_bytes(archive)
    import zipfile
    with zipfile.ZipFile(io.BytesIO(archive)) as zipped:
        (output / "delivery_manifest.json").write_bytes(zipped.read("delivery_manifest.json"))
    return {"directory": str(output), "zip": str(archive_path), "zip_bytes": len(archive), **payload["summary"]}


_README = """NanoCoFold · 신규 VHH 설계 · hotspot 사용/미사용 · 전달용 보고서

ZIP 전체를 풀고 report.html을 최근 Chrome·Edge·Firefox에서 여세요.
인터넷·Python·GPU 설치 없이 표 정렬, MolStar 구조 탐색, PAE 확인과 원본 구조 저장을 사용할 수 있습니다.
WebGL과 브라우저의 gzip DecompressionStream 지원이 필요합니다. PAE 기능을 지원하지 않는 브라우저에서도 표와 구조는 남습니다.
표적·생성 후보·예측 모델은 서로 다른 기록입니다. 예상 수, 실제 생성 수, 실행/미시작/실패/평가 수를 구분합니다.
전체 저장 모델을 포함하며 화면 필터는 JSON/CSV/XLSX/FASTA 원본 범위를 바꾸지 않습니다.
Excel 셀 한도를 넘는 긴 출처 기록은 XLSX의 Long fields 시트에 순서대로 분할합니다. 본문 셀에 분할 행과 원문 SHA256을 표시하며, JSON/CSV에는 원문을 그대로 보존합니다.
표의 기본 ipSAE 정렬은 탐색 순서입니다. 엔진 간 합산 순위·결합 확률·KD·실험 검증이 아닙니다.
구조 검토 자료가 포함되면 기본 표는 같은 표적·조건 내 후보 검토 순위로 정렬합니다(동률 허용).
품질·지정 위치·독립 엔진의 자세 일치·생성 단계 구조 일관성을 따로 확인합니다. 1위도 보류일 수 있습니다.
검토 정책 v2는 같은 자세에서 두 엔진 이상이 품질·지정 위치를 충족한 다중 엔진 지지와 엔진별 단독 유망을 구분합니다. 한 엔진만 유망해도 합의 부족을 실패로 간주하지 않습니다.
엔진 간 자세 차이와 생성·재예측 불일치는 별도 진단이며 자동 비결합·탈락 사유가 아닙니다. 이전 정책의 판정은 pose_assessment.json의 버전과 함께 읽으세요.
검토 임계값은 탐색용이며 실험 결합에 대해 보정되지 않았습니다. 원본 점수·기록 적격·엔진별 순위는 보존합니다.
새 후보 서열에는 native DockQ를 성공 기준으로 적용하지 않습니다. hotspot 사용 조건은 알려진 계면의 위치 정보를 제공받은 설계입니다.
동일한 요청 seed만으로 두 생성기가 같은 난수 제어를 제공한다고 추정하지 않습니다. BoltzGen의 upstream seed 지원 여부는 조건별 generation_engine_info 기록을 확인하세요.
좌표는 원본 UTF-8 bytes 그대로 한 번씩 내장합니다. PAE는 float64 배열을 반올림 없이 gzip으로 보관합니다.
검증된 PDB 참조가 제공되면 회색으로 함께 표시하며, 표시를 켜고 끌 수 있습니다. 파랑은 예측 표적, 금색은 새 VHH입니다.
기본은 리본 구조입니다. 구조 위에서 서열 패널과 잔기·곁사슬을 각각 켜거나 끌 수 있고 선택 해제로 초록색 강조를 지울 수 있습니다.
참조 복합체는 표적 Cα만으로 예측 표적에 정렬합니다. 참조 VHH는 정렬에 사용하지 않으며 새 VHH와 서열이 다릅니다. 겹침은 결합 위치 비교용이며 결합 검증이 아닙니다.
PDB 참조는 추론 완료 후 보고서에만 추가합니다. 원본 예측 좌표·점수는 그대로 보존하며 원본 예측 구조 저장에는 참조·정렬 변환이 포함되지 않습니다.
hotspot_assessments.csv는 기존 접촉 수·판정을 설명합니다. hotspot 미지정·미평가·번호 대응 실패는 0접촉과 구분합니다.
reference_binding_sites.csv와 XLSX의 Reference binding sites 시트는 표적–VHH 단백질 중원자 5 Å 접촉 잔기의 겹침을 기록합니다. 양쪽에서 좌표가 있는 표적 위치만 비교합니다. 회복률=공통 접촉 잔기/참조 접촉 잔기, Jaccard=교집합/합집합이며, 결합·KD·native DockQ 또는 같은 자세를 입증하지 않습니다.
reference_alignment.csv와 XLSX의 Visual target alignment 시트는 시각화 정렬의 출처·변환·대응 표적 위치를 기록합니다. 이는 모델 선별 점수와 별개입니다.
MSA·가중치·로그·중간 계산·중복 개별 HTML은 포함하지 않습니다. 재실행 또는 전체 감사용 원본은 별도로 보관하세요.
현재 run의 API 요청과 이전 MSA cache ticket의 누적 요청은 cases.csv의 별도 열입니다.
delivery_manifest.json은 전달 파일 bytes/해시, source 파일과 내장 좌표/PAE의 출처를 기록합니다.
OpenDDE 추가 전달본은 기존 세 엔진과 추가 계산의 source_campaigns 및 timing_scopes를 구분합니다. timing_scopes.csv에 각 계산의 원래 시작·종료·경과 시간을 보존하며, 이를 네 엔진의 새 단일 wall time으로 합치지 않습니다.
"""


_DOCUMENT = r'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>NanoCoFold · 신규 VHH hotspot 비교</title><style>__MOLSTAR_CSS__</style><style>
:root{font-family:system-ui,-apple-system,sans-serif;color:#20343d;background:#f5f7f8;font-size:18px;line-height:1.6}*{box-sizing:border-box}body{margin:0}main{max-width:1480px;margin:auto;padding:32px}h1{font-size:32px;line-height:1.3}h2{font-size:25px;margin:30px 0 16px}h3{font-size:20px}.muted{color:#536771}a{color:#145d70}button,select,input{font:inherit;min-height:48px;border:1px solid #b9cbd3;border-radius:4px;padding:8px 12px;background:white;color:inherit}button{cursor:pointer}button:hover{background:#e5eff3}button:focus-visible,select:focus-visible,input:focus-visible{outline:3px solid #146e7d;outline-offset:2px}.sheet{padding:24px;background:white;border:1px solid #d6e1e6;border-radius:6px;margin:20px 0}.stats{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:16px}.stat{border-left:3px solid #5d8b99;padding:12px 16px;background:#edf3f6}.stat strong{display:block;font-size:27px}.controls{display:flex;flex-wrap:wrap;gap:16px;align-items:end}.controls label{display:flex;flex-direction:column;gap:6px;min-width:180px}table{border-collapse:collapse;width:100%;font-size:16px}th,td{border-bottom:1px solid #dce5ea;padding:12px;text-align:left;vertical-align:top}th{background:#eef4f6}#model-table{min-width:1650px}#model-table th{white-space:nowrap}#model-table td:last-child{min-width:240px;max-width:360px;overflow-wrap:anywhere}th button{min-height:48px;border:0;padding:0;background:transparent;text-align:left;font-weight:700}tbody tr.selected{background:#e5f1f5}.site-assessment{border-left:4px solid #536771;background:#f0f4f6;padding:16px;margin:20px 0}.site-assessment[data-status="no_contact"],.site-assessment[data-status="threshold_unmet"],.site-assessment[data-status="mapping_failed"]{border-color:#9b4c27;background:#fff4e9}.site-assessment p{margin:8px 0}.site-warning{color:#85421e;font-weight:600}td button{white-space:nowrap}.scroll{overflow:auto;max-height:700px}.plots{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:18px;margin-bottom:22px}.plot{height:310px}.plot-panel{border:1px solid #d6e1e6;padding:12px;min-width:0}#viewer{height:600px;position:relative;border:1px solid #c5d5dd;margin:18px 0}#viewer .msp-plugin button,#viewer .msp-plugin select,#viewer .msp-plugin input{min-height:initial}#viewer .msp-viewport-controls-panel{max-width:var(--menu-width,290px)}#viewer .msp-viewport-controls-panel-controls{max-height:var(--menu-height,400px);overflow:auto}.status{padding:12px 0;min-height:48px;overflow-wrap:anywhere}.metric-row{display:flex;flex-wrap:wrap;gap:16px}.metric-row>div{min-width:140px;padding:12px;background:#eff4f6}.metric-row strong{display:block;font-size:24px}details{margin:14px 0}summary{cursor:pointer}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:16px}.downloads{display:flex;flex-wrap:wrap;gap:16px}nav{display:flex;flex-wrap:wrap;gap:18px;border-bottom:1px solid #dce5ea;padding:12px 0}.structure-legend{display:flex;flex-wrap:wrap;gap:12px 24px;margin-top:20px}.structure-legend span{display:flex;align-items:center;gap:8px}.structure-legend i{width:16px;height:16px;border-radius:50%;display:inline-block}.reference-toggle{display:flex;align-items:center;gap:12px;min-height:48px;cursor:pointer}.reference-toggle input{width:22px;height:22px;min-height:22px;padding:0}#case-failures{border-left:3px solid #a15c30;padding:0 16px}@media(max-width:1000px){.plots{grid-template-columns:1fr}.stats{grid-template-columns:1fr 1fr}main{padding:20px}#viewer{height:470px}}@media(max-width:700px){main{padding:14px}.sheet{padding:16px}h1{font-size:28px}.controls label{width:100%}#viewer{height:400px}#viewer .msp-sequence-select{display:flex;flex-wrap:wrap;height:48px}#viewer .msp-sequence-wrapper-non-empty{top:49px}.stats{gap:8px}}@media print{.controls,.plot,#viewer,.downloads{display:none}.scroll{max-height:none;overflow:visible}main{padding:0}.sheet{border:0;padding:0}table{font-size:12px}h1{font-size:24px}}

.sheet{min-width:0;overflow-wrap:anywhere}.scroll{min-width:0;width:100%;max-width:100%}.controls label{max-width:100%}.controls input,.controls select{max-width:100%}
#viewer .msp-plugin .msp-sequence-select{height:64px!important;margin-bottom:1px;display:flex;align-items:center;flex-wrap:nowrap;overflow-x:auto;overflow-y:hidden;white-space:nowrap}
#viewer .msp-plugin .msp-sequence-select>span{flex:0 0 auto;font-size:16px!important;line-height:40px!important;padding:0 10px}
#viewer .msp-plugin .msp-sequence-select>select{flex:0 0 auto;max-width:160px;height:40px!important;min-height:40px!important;line-height:40px!important;font-size:16px!important;margin:0;padding:0 22px 0 8px}
#viewer .msp-plugin .msp-sequence-wrapper-non-empty{top:65px!important;font-size:16px!important;line-height:1.8!important;padding:22px 10px 8px}
#viewer .msp-plugin .msp-sequence-number{font-size:12px!important;line-height:1.2!important}
#viewer .msp-plugin .msp-layout-top{height:160px!important}
#viewer .msp-plugin .msp-layout-main{top:160px!important}
#viewer .msp-plugin .msp-layout-hide-top .msp-layout-main{top:0!important}
.viewer-display-toolbar{display:flex;align-items:center;flex-wrap:wrap;gap:16px;padding:12px 16px;border:1px solid #d6e1e6;border-radius:6px;background:#f5f7fb;margin:16px 0}
.viewer-display-toolbar label{display:flex;align-items:center;gap:8px;cursor:pointer}
.viewer-display-toolbar input{width:18px;height:18px;min-height:18px;padding:0;margin:0}
.viewer-display-toolbar button{margin-left:auto}
</style></head><body><main><p class="muted">NanoCoFold · 저장된 결과 · 전달용 보고서</p><h1>신규 VHH 설계 · hotspot 사용/미사용 비교</h1><p id="engine-description"></p><nav><a href="#counts">01 실행 요약</a><a href="#timings">02 소요 시간</a><a href="#models">03 후보 비교</a><a href="#structures">04 구조·PAE</a><a href="#provenance">05 출처·전달</a></nav>
<section class="sheet" id="counts"><h2>01 실행 요약</h2><div class="stats" id="stats"></div><p id="run-note"></p><p class="muted">예측 신뢰도는 실제 결합·친화도·KD 또는 실험 검증을 뜻하지 않습니다. hotspot 사용 설계에는 알려진 계면 위치 정보가 제공됩니다. 신규 후보의 서열은 정답 VHH와 다르므로 native DockQ를 성공 기준으로 적용하지 않았습니다.</p><div id="case-failures"></div><details><summary>표적·정확한 입력 서열</summary><div id="targets"></div></details><details><summary>생성 기록과 평가 제한</summary><div id="generation"></div></details></section>
<section class="sheet" id="timings"><h2>02 NanoCoFold 소요 시간</h2><p id="time-note"></p><div class="scroll"><table id="timing-scopes"></table></div><div id="timing-definitions"></div><div class="scroll"><table id="timing-table"></table></div><details><summary>조건·단계별 실측 시간 요약</summary><div class="scroll"><table id="timing-summary"></table></div></details><p class="muted">시간은 관측된 wall time입니다. 추론 준비·I/O·대기를 포함할 수 있으며 순수 GPU kernel 시간이 아닙니다. cache 이전 획득 시간과 이번 네트워크 요청 수는 분리했습니다. 단계별 시간에는 중복 범위가 있으므로 합계를 전체 wall time으로 해석하지 않습니다.</p></section>
<section class="sheet" id="models"><h2>03 후보 비교</h2><div class="controls"><label>표적<select id="target-filter"></select></label><label>hotspot<select id="arm-filter"><option value="all">전체 조건</option><option value="off">미사용</option><option value="on">사용</option></select></label><label>예측 엔진<select id="engine-filter"><option value="all">전체 엔진</option></select></label><label>검색<input id="search" placeholder="후보·서열·오류"></label></div><p id="model-count"></p><p id="sort-note" class="muted"></p><p class="muted">신뢰도 점수와 결합 위치 일치는 별개입니다. hotspot 접촉은 지정 잔기의 충족 여부이며, 참조 접촉 잔기 회복률은 알려진 표적 결합 부위와의 겹침입니다. PAE는 참조 구조와의 실측 오차가 아니며 pDockQ2는 실측 DockQ가 아닙니다. hotspot 사용은 후보 생성 조건이며 후속 위치 제약 적용 여부는 선택 모델에서 확인하세요.</p><div class="scroll"><table id="model-table"></table></div></section>
<section class="sheet" id="structures"><h2>04 선택 모델 · 구조·PAE</h2><h3 id="selected-title">모델을 선택하세요</h3><h3>예측 계면 신뢰도</h3><div class="metric-row" id="metrics"></div><div id="binding-site-assessment" class="site-assessment" role="status" aria-live="polite"></div><details><summary>결합 위치 진단 · 판정값·임계값·참조 접촉 잔기</summary><pre id="binding-site-record"></pre></details><div class="plots"><div class="plot-panel"><h3>seed별 기록</h3><div id="seed-plot" class="plot"></div><p class="muted">같은 후보의 저장 관측을 표시합니다. 단일 seed는 재현성 검증이 아닙니다.</p></div><div class="plot-panel"><h3>PAE · 예측 정렬 오차 (Å)</h3><div id="pae-plot" class="plot"></div><p id="pae-note" class="muted"></p></div><div class="plot-panel"><h3>계면 신뢰도</h3><div id="interface-plot" class="plot"></div><p class="muted">동일 표적·조건의 모델을 표시합니다. 엔진 간 점수를 합산하지 않습니다.</p></div></div><div class="structure-legend" aria-label="구조 색상"><span><i style="background:#3b82b2"></i>예측 표적</span><span><i style="background:#c49a40"></i>새 VHH</span><span><i style="background:#697582"></i>PDB 참조 복합체</span></div><label class="reference-toggle"><input id="reference-toggle" type="checkbox" checked> 회색 PDB 참조 겹쳐 보기</label><p id="reference-status" class="muted status" role="status" aria-live="polite"></p><p class="muted">참조는 표적 Cα만으로 정렬합니다. 새 VHH와 기존 VHH의 서열은 다르므로 결합 위치를 비교하는 용도입니다. 원본 예측 좌표·점수는 바뀌지 않습니다.</p><p id="viewer-status" class="status" role="status" aria-live="polite">저장 구조를 선택하세요.</p><div class="viewer-display-toolbar" aria-label="구조 표시 옵션"><label><input id="sequence-toggle" type="checkbox"> 서열 패널 표시</label><label><input id="residues-toggle" type="checkbox"> 잔기·곁사슬 표시</label><button id="clear-selection" type="button">선택 해제</button></div><div id="viewer" aria-label="선택된 원본 예측 구조 MolStar"></div><div class="controls"><button id="fit" type="button">전체 보기</button><button id="target-focus" type="button">표적</button><button id="binder-focus" type="button">VHH</button><button id="download-structure" type="button">원본 예측 구조 저장</button></div><details><summary>선택 후보의 정확한 서열·원본 모델 기록</summary><pre id="selected-record"></pre></details></section>
<section class="sheet" id="provenance"><h2>05 출처·전달</h2><div class="downloads"><a href="benchmark.xlsx">전체 점수·시간 XLSX</a><a href="models.csv">모델 CSV</a><a href="candidate_review.csv">후보 검토 순위 CSV</a><a href="pose_pairs.csv">모델 간 자세 비교 CSV</a><a href="pose_assessment.json">구조 판정·임계값 JSON</a><a href="generated_candidates.csv">생성 기록 CSV</a><a href="cases.csv">조건·시간 CSV</a><a href="candidate_sequences.fasta">후보 FASTA</a><a href="target_sequences.fasta">표적 FASTA</a><a href="benchmark.json">출처 JSON</a><a href="hotspot_assessments.csv">hotspot 진단 CSV</a><a href="reference_binding_sites.csv">참조 결합 부위 CSV</a><a href="delivery_manifest.json">전달 파일 해시</a></div><p>표·JSON·CSV·XLSX는 같은 전체 모델 기록을 사용합니다. 화면 필터는 원본 내보내기 범위를 변경하지 않습니다. MSA·가중치·로그·중간 계산 파일은 전달본에서 제외했습니다.</p><p class="muted">ipSAE·ipTM는 예측 배치의 신뢰도입니다. 낮은 PAE는 낮은 예측 정렬 오차를 뜻합니다. pDockQ2는 AlphaFold-Multimer에서 보정된 지표이며 Boltz-2·Chai-1·OpenDDE 값은 탐색용입니다. OpenDDE는 Preview 엔진이며, 공통 선별 기준 통과는 실제 결합이나 엔진별 성능 보정을 입증하지 않습니다. 원본 score·서열·좌표는 수정하지 않았습니다.</p><details><summary>실패·미평가·출처 기록</summary><pre id="source-record"></pre></details></section></main>
<script id="benchmark-data" type="application/json">__BENCHMARK_DATA__</script><script>__PLOTLY_JS__</script><script>__MOLSTAR_JS__</script><script>__BENCHMARK_VIEWER_DISPLAY_JS__</script><script>
(() => {"use strict";
const data=JSON.parse(document.getElementById('benchmark-data').textContent),$=id=>document.getElementById(id),names={colabfold:'ColabFold',boltz2:'Boltz-2',chai1:'Chai-1',opendde:'OpenDDE Preview'},colors={colabfold:'#326f9a',boltz2:'#428580',chai1:'#8064a0',opendde:'#a86631'},engines=data.engines||['colabfold','boltz2','chai1'],scopeNames={original_3engine:'기존 세 엔진 계산',opendde_addition:'OpenDDE 추가 계산'};
const fmt=value=>typeof value==='number'&&Number.isFinite(value)?(Number.isInteger(value)?String(value):value.toFixed(3)):value===true?'통과':value===false?'미통과':value===null||value===undefined?'미평가':String(value);
const el=(tag,text)=>{const n=document.createElement(tag);if(text!==undefined)n.textContent=text;return n;};
engines.forEach(engine=>$('engine-filter').append(new Option(names[engine]||engine,engine)));
$('engine-description').textContent=`RFantibody·BoltzGen의 후보와 ${engines.map(engine=>names[engine]||engine).join('·')}의 저장 평가 기록을 비교합니다. 예상 규모와 실제 생성·실패·미평가를 구분합니다.`;
if(data.timing_scopes){const link=el('a','계산 범위·시간 CSV');link.href='timing_scopes.csv';document.querySelector('.downloads').append(link);}
let sort=data.pose_assessment?'review_rank':'ipsae',descending=!data.pose_assessment,selected=null,viewerPromise,queue=Promise.resolve(),serial=0,currentStructure=null,currentReference=null;
const s=data.summary;for(const [label,value,plan]of [['생성 완료 후보',s.generated_complete,s.expected_candidates],['시도한 예측',s.models_attempted,s.expected_models],['지표 평가 완료',s.models_evaluated,s.expected_models],['기록 적격 모델',s.models_eligible,null]]){const n=el('div');n.className='stat';n.append(el('span',label),el('strong',String(value)+(plan===null?'':' / 예상 '+plan)));$('stats').append(n);}
$('run-note').textContent=`실행 상태 ${s.status||'미기록'} · 조건 기록 ${s.recorded_cases}/${s.expected_cases??'미기록'} · 고유 후보 서열 ${s.unique_sequences} · 실패 ${s.models_failed} · 중단 ${s.models_interrupted||0} · 실행중 ${s.models_running||0} · 예측 미시작 ${s.models_not_started}. 후보 기록과 고유 서열 수는 다를 수 있습니다.`+(data.timing_scopes?` 기존 세 엔진 ${s.legacy_models_preserved}개 모델을 보존하고 OpenDDE ${s.added_model_records}개 기록을 추가했습니다. 후보 재생성은 0건입니다.`:'');
for(const c of data.cases.filter(c=>c.status!=='complete'||c.error||c.issues.length)){const p=el('p',`${c.pdb_id} · hotspot ${c.arm==='on'?'사용':'미사용'} · ${scopeNames[c.source_scope]||''} · ${c.status} · ${c.error||c.issues.join(' / ')||'일부 자료 미평가'}`);$('case-failures').append(p);}
for(const t of data.targets){const d=el('details');d.append(el('summary',`${t.pdb_id} · ${t.name||''} · ${t.sequence_length} aa`),el('pre',t.sequence+'\nSHA256 '+t.sequence_sha256));if(typeof t.source_url==='string'&&/^https?:\/\//i.test(t.source_url)){const a=el('a','원본 표적 출처');a.href=t.source_url;d.append(a);}$('targets').append(d);}
$('generation').append(el('pre',JSON.stringify({generated_records:s.generated_records,successful_candidates:s.generated_complete,unique_sequences:s.unique_sequences,records:data.generated_candidates},null,2)));
$('time-note').textContent=data.timing_scopes?'기존 세 엔진 계산과 OpenDDE 추가 계산의 시작·종료·경과 시간을 각각 보존합니다. 합친 보고서는 네 엔진을 새로 계산한 하나의 wall time을 제시하지 않습니다.':`전체 경과 ${s.elapsed_seconds===null?'미기록':fmt(s.elapsed_seconds)+'초'} · 시작 UTC ${s.started_utc||'미기록'} · 종료 UTC ${s.completed_utc||'미기록'}`;
const stages=[...new Set(data.cases.flatMap(c=>Object.keys(c.timings||{})))],scopeKeys=data.timing_scopes?['source_scope']:[];const timeKeys=[...scopeKeys,'pdb_id','arm','status','elapsed_seconds',...stages,'models_evaluated','network_requests_this_run','source_ticket_lifetime_requests'];const timeRows=data.cases.map(c=>({...c,...c.timings}));
function plainTable(host,rows,keys){host.replaceChildren();const head=el('thead'),tr=el('tr');keys.forEach(k=>tr.append(el('th',k==='source_scope'?'계산 범위':k)));head.append(tr);const body=el('tbody');for(const r of rows){const line=el('tr');keys.forEach(k=>line.append(el('td',fmt(k==='source_scope'?(scopeNames[r[k]]||r[k]):r[k]))));body.append(line);}host.append(head,body);}plainTable($('timing-table'),timeRows,timeKeys);plainTable($('timing-summary'),data.timing_summary,[...scopeKeys,'arm','stage','observations','total_seconds','median_seconds','min_seconds','max_seconds']);if(data.timing_scopes)plainTable($('timing-scopes'),data.timing_scopes,['source_scope','status','elapsed_seconds','started_utc','completed_utc','elapsed_scope','timestamp_scope']);for(const timing of data.timing_scopes||[s]){if(timing.timing_note)$('timing-definitions').append(el('p',`${scopeNames[timing.source_scope]||'전체 실행'} · ${timing.timing_note}`));}
$('target-filter').append(new Option('전체 표적','all'));data.targets.forEach(t=>$('target-filter').append(new Option(t.pdb_id+' · '+(t.name||''),t.pdb_id)));
const statusNames={complete:'예측 완료',failed:'실패',not_started:'미시작',running:'실행중',partial_failure:'부분 실패',interrupted:'중단',cancelled:'취소'},reasonNames={clash:'원자 충돌',ipsae:'ipSAE 기준 미달',iptm:'ipTM 기준 미달',interface_plddt:'계면 pLDDT 기준 미달',developability:'서열 위험 기준 미달',developability_unassessed:'CDR 서열 위험 미평가',hotspot:'hotspot 접촉 기준 미달',scoring_error:'지표 계산 오류'},captions={review_rank:'검토 순위 (동률 허용)',candidate_review_priority:'후보 검토 판정',review_priority:'모델 검토 판정',pose_quality_supporting_engines:'품질·위치 지지 엔진 수',pose_supporting_engines:'자세 일치 엔진 수',pose_consensus_status:'모델 자세 일치',design_consistency_status:'생성 구조 일관성',design_binder_rmsd_A:'생성 대비 VHH RMSD (Å)',pdb_id:'표적',arm:'hotspot',candidate_label:'후보',design_origin:'생성 도구',predictor_engine:'예측 도구',seed:'seed',ipsae:'ipSAE',iptm:'ipTM',interface_pae:'계면 PAE (Å)',interface_plddt:'계면 pLDDT',pdockq2:'pDockQ2',hotspot_contact_frac:'hotspot 접촉률',hotspot_site_status:'지정 위치 충족',reference_contact_recovery:'참조 접촉 잔기 회복률',final_eligible:'기록 적격',prediction_status:'예측 상태',elapsed_seconds:'예측 시간 (초)',error:'평가·제외 이유'};
const poseNames={multi_engine_supported:'다중 엔진 지지·탐색용',single_engine_promising:'엔진별 단독 유망·탐색용',prioritized:'우선 검토',site_unmet:'보류: 지정 위치 미충족',site_not_assessed:'보류: 위치 미평가',confidence_gate_unmet:'보류: 품질 기준 미달',confidence_not_assessed:'보류: 품질 미평가',geometry_not_assessed:'보류: 구조 미평가',design_inconsistent:'보류: 생성 구조 불일치',design_not_assessed:'보류: 생성 구조 미평가',consensus_insufficient:'보류: 독립 엔진 지지 부족',supported:'품질·위치·자세 지지',quality_support_insufficient:'자세 일치·품질 지지 부족',single_engine:'단일 엔진',discordant:'자세 불일치',consistent:'일관성 기준 충족',inconsistent:'일관성 기준 미달',not_applicable:'해당 없음',not_assessed:'미평가',disabled:'검토 끔'};const poseLabel=v=>v==null?'미평가':(poseNames[v]||'미평가');
const diagnosticValue=(r,k)=>{const h=data.hotspot_assessments?.[r.id];return k==='hotspot_contact_frac'?h?.contact_fraction:k==='hotspot_site_status'?(h&&!['not_assessed','not_requested','mapping_failed'].includes(h.status)?h.label:null):k==='reference_contact_recovery'?data.reference_binding_sites?.[r.id]?.reference_contact_recovery:r[k];};
const keys=['pdb_id','arm','review_rank','candidate_label','candidate_review_priority','review_priority','pose_quality_supporting_engines','pose_supporting_engines','pose_consensus_status','design_consistency_status','design_binder_rmsd_A','predictor_engine','hotspot_contact_frac','hotspot_site_status','reference_contact_recovery','ipsae','iptm','interface_pae','interface_plddt','pdockq2','final_eligible','prediction_status','elapsed_seconds','design_origin','seed','error'];function filtered(){const q=$('search').value.toLowerCase();return data.models.filter(r=>($('target-filter').value==='all'||r.pdb_id===$('target-filter').value)&&($('arm-filter').value==='all'||r.arm===$('arm-filter').value)&&($('engine-filter').value==='all'||r.predictor_engine===$('engine-filter').value)&&(!q||[r.id,r.sequence,r.error,r.candidate_label].join(' ').toLowerCase().includes(q))).sort((a,b)=>{const x=diagnosticValue(a,sort),y=diagnosticValue(b,sort);if(x==null)return y==null?a.id.localeCompare(b.id):1;if(y==null)return-1;const cmp=typeof x==='number'&&typeof y==='number'?x-y:String(x).localeCompare(String(y));return(descending?-cmp:cmp)||a.id.localeCompare(b.id);});}
function render(){const rows=filtered();$('model-count').textContent=`표시 ${rows.length} / 전체 ${data.models.length} 모델 기록`;$('sort-note').textContent=`정렬 ${captions[sort]||sort} · ${descending?'높은 값부터':'낮은 값부터'} · 미평가 마지막 · 검토 순위는 같은 표적·조건 내 동률 순위이며, 1위도 보류일 수 있습니다. 점수 합산·결합 확률 순위가 아닙니다.`;const h=el('thead'),hr=el('tr'),body=el('tbody');for(const k of keys){const th=el('th'),b=el('button',(captions[k]||k)+(sort===k?(descending?' ↓':' ↑'):''));b.type='button';b.addEventListener('click',()=>{descending=sort===k?!descending:!['pdb_id','arm','candidate_label','design_origin','predictor_engine','prediction_status','hotspot_site_status','error','review_rank','design_binder_rmsd_A','candidate_review_priority','review_priority','pose_consensus_status','design_consistency_status'].includes(k);sort=k;render();});th.setAttribute('aria-sort',sort===k?(descending?'descending':'ascending'):'none');th.append(b);hr.append(th);}h.append(hr);for(const r of rows){const tr=el('tr');tr.dataset.id=r.id;tr.classList.toggle('selected',selected?.id===r.id);for(const k of keys){const td=el('td');td.dataset.field=k;const v=diagnosticValue(r,k);td.dataset.value=v==null?'null':String(v);if(k==='candidate_label'){const b=el('button',r.candidate_label);b.type='button';b.title=r.id;b.addEventListener('click',()=>{choose(r);$('structures').scrollIntoView({block:'start',behavior:'smooth'});});td.append(b);}else if(['candidate_review_priority','review_priority','pose_consensus_status','design_consistency_status'].includes(k)){td.textContent=poseLabel(v);td.title=r.review_priority_reason||'탐색용 구조 검토 판정';}else if(k==='review_rank'){td.textContent=v==null?'미평가':String(v);td.title='같은 표적·조건 내 후보 검토 순위(동률 허용). 1위도 보류일 수 있습니다.';}else if(k==='hotspot_contact_frac'){const h=data.hotspot_assessments?.[r.id];td.textContent=h?.contact_fraction==null?(h?.status==='not_requested'?'미지정':'미평가'):`${h.contacted}/${h.requested} · ${(100*h.contact_fraction).toFixed(1)}%`;td.title='지정 hotspot 접촉 수 / 요청 수. 기존 접촉률·임계값을 보존합니다.';}else if(k==='hotspot_site_status'){td.textContent=data.hotspot_assessments?.[r.id]?.label||'미평가';if(['no_contact','threshold_unmet','mapping_failed'].includes(data.hotspot_assessments?.[r.id]?.status))td.className='site-warning';}else if(k==='reference_contact_recovery'){const f=data.reference_binding_sites?.[r.id];td.textContent=f?.available?`${(100*v).toFixed(1)}% · ${f.shared_contacts}/${f.reference_contacts}`:'비교 불가';td.title='표적–VHH 중원자 5 Å 접촉. 참조 접촉 표적 잔기 중 예측에서도 접촉하는 비율. 다른 VHH 서열의 결합 부위 겹침이며 같은 자세·실제 결합을 입증하지 않습니다.';}else td.textContent=k==='arm'?(r.arm==='on'?'사용':'미사용'):k==='predictor_engine'?(names[r[k]]||r[k]):k==='prediction_status'?(statusNames[r[k]]||r[k]):k==='error'?(r.error||(r.exclusion_reason?String(r.exclusion_reason).split(';').filter(Boolean).map(code=>reasonNames[code]||code).join(' · '):r.final_eligible===true?'기준 통과':r.evaluated?'판정 이유 미기록':'미평가')):fmt(r[k]);tr.append(td);}body.append(tr);}$('model-table').replaceChildren(h,body);}
['target-filter','arm-filter','engine-filter','search'].forEach(id=>$(id).addEventListener(id==='search'?'input':'change',render));
async function paePlot(r,version){const entry=data.pae_arrays[r.pae_key];if(!entry){Plotly.purge('pae-plot');$('pae-note').textContent='PAE 배열 미평가 또는 파일 없음';return;}try{if(typeof DecompressionStream==='undefined')throw Error('이 브라우저의 gzip 배열 해제 미지원');const zipped=Uint8Array.from(atob(entry.data),c=>c.charCodeAt(0));const stream=new Blob([zipped]).stream().pipeThrough(new DecompressionStream('gzip'));const bytes=await new Response(stream).arrayBuffer();if(bytes.byteLength!==entry.bytes)throw Error('PAE 내장 배열 크기 불일치');const values=new DataView(bytes),z=[];for(let i=0;i<entry.dimension;i++){const row=[];for(let j=0;j<entry.dimension;j++)row.push(values.getFloat64((i*entry.dimension+j)*8,true));z.push(row);}if(version!==serial)return;Plotly.newPlot('pae-plot',[{type:'heatmap',z,x:r.pae_labels||entry.labels,y:r.pae_labels||entry.labels,zmin:0,zmax:30,colorscale:'Blues',colorbar:{title:'Å'},hovertemplate:'행 %{y}<br>열 %{x}<br>PAE %{z:.3f} Å<extra></extra>'}],{margin:{l:50,r:20,t:10,b:50},xaxis:{title:'열 잔기'},yaxis:{title:'행 잔기',autorange:'reversed'}},{responsive:true,displayModeBar:false});$('pae-note').textContent='행·열은 원본 예측 토큰 순서입니다. 낮은 값이 작은 상대 배치 오차를 뜻합니다.';}catch(error){Plotly.purge('pae-plot');$('pae-note').textContent='PAE 표시 제한: '+error.message;}}
function otherPlots(r){const same=data.models.filter(m=>m.case_id===r.case_id&&m.source_sequence_id===r.source_sequence_id),cohort=data.models.filter(m=>m.case_id===r.case_id);const seed=[],inter=[];for(const e of engines){const obs=same.filter(m=>m.predictor_engine===e&&Number.isFinite(m.iptm)&&Number.isFinite(m.seed));seed.push({type:'scatter',mode:'markers',name:names[e],x:obs.map(m=>m.seed),y:obs.map(m=>m.iptm),marker:{color:colors[e],size:10},text:obs.map(m=>m.candidate_label)});const group=cohort.filter(m=>m.predictor_engine===e&&Number.isFinite(m.iptm)&&Number.isFinite(m.ipsae));inter.push({type:'scatter',mode:'markers',name:names[e],x:group.map(m=>m.iptm),y:group.map(m=>m.ipsae),text:group.map(m=>m.candidate_label),marker:{color:colors[e],size:10},hovertemplate:'%{text}<br>ipTM %{x:.3f}<br>ipSAE %{y:.3f}<extra>%{fullData.name}</extra>'});}const layout={margin:{l:50,r:10,t:10,b:60},legend:{orientation:'h',y:-.25}};Plotly.newPlot('seed-plot',seed,{...layout,xaxis:{title:'seed'},yaxis:{title:'ipTM',range:[0,1]}},{responsive:true,displayModeBar:false});Plotly.newPlot('interface-plot',inter,{...layout,xaxis:{title:'ipTM',range:[0,1]},yaxis:{title:'ipSAE',range:[0,1]}},{responsive:true,displayModeBar:false});}
const host=$('viewer');const resizeMenus=()=>{host.style.setProperty('--menu-width',Math.max(120,host.clientWidth-68)+'px');host.style.setProperty('--menu-height',Math.max(160,host.clientHeight-80)+'px');};resizeMenus();new ResizeObserver(resizeMenus).observe(host);
const getViewer=()=>viewerPromise ||= molstar.Viewer.create('viewer',{extensions:[],layoutIsExpanded:false,layoutShowControls:true,layoutShowRemoteState:false,layoutShowSequence:true,layoutShowLog:false,layoutShowLeftPanel:false,collapseRightPanel:true,viewportShowExpand:false,viewportShowToggleFullscreen:false,viewportShowControls:true,viewportShowSelectionMode:true,viewportFocusBehavior:'disabled',volumeStreamingDisabled:true,viewportBackgroundColor:'#ffffff'});

let currentDisplay=null;
const displayOptions=()=>({sequence:$('sequence-toggle').checked,residues:$('residues-toggle').checked});
function alignedReferencePdb(asset,fit){
  return asset.content.split('\n').map(line=>{
    if(!/^(ATOM  |HETATM)/.test(line))return line;
    const xyz=[line.slice(30,38),line.slice(38,46),line.slice(46,54)].map(Number);
    if(!xyz.every(Number.isFinite))throw Error('참조 좌표 형식 오류');
    const mapped=fit.translation.map((t,j)=>t+xyz.reduce((sum,v,i)=>sum+v*fit.rotation[i][j],0));
    const fields=mapped.map(v=>{if(!Number.isFinite(v))throw Error('참조 정렬 좌표 오류');const s=v.toFixed(3);if(s.length>8)throw Error('참조 좌표 표시 범위 초과');return s.padStart(8,' ');});
    return line.slice(0,30)+fields.join('')+line.slice(54);
  }).join('\n');
}
function referenceStatus(r){
  const fit=data.reference_overlays?.[r.id];
  if(!fit?.available)return 'PDB 참조 겹침 없음 · '+(fit?.reason||'보고서에 참조 구조가 포함되지 않았습니다.');
  return `PDB ${r.pdb_id} 참조 · ${$('reference-toggle').checked?'회색 표시':'표시 끔'} · 표적 Cα ${fit.matched_ca}개 정렬 · 정렬 RMSD ${fit.rmsd_angstrom.toFixed(3)} Å (시각화 정렬값)`;
}
function bindingSite(r){
  const h=data.hotspot_assessments?.[r.id],f=data.reference_binding_sites?.[r.id],host=$('binding-site-assessment');
  host.replaceChildren();host.dataset.status=h?.status||'not_assessed';host.append(el('h3','결합 위치 · 신뢰도와 별도 확인'));
  host.append(el('h3',`후보 검토 순위 ${fmt(r.review_rank)} · ${poseLabel(r.candidate_review_priority)}`));
  host.append(el('p',`모델 검토 판정: ${poseLabel(r.review_priority)}`));
  host.append(el('p',`품질·위치 지지 엔진 ${fmt(r.pose_quality_supporting_engines)}개 · 자세 일치 엔진 ${fmt(r.pose_supporting_engines)}개 · 모델 판정 ${poseLabel(r.pose_consensus_status)}`));
  host.append(el('p',`생성 단계 구조 일관성: ${poseLabel(r.design_consistency_status)} · 표적 정렬 후 VHH RMSD ${fmt(r.design_binder_rmsd_A)} Å`));
  host.append(el('p','검토 정책 v2는 품질·지정 위치를 충족한 엔진별 단독 유망과 같은 자세의 다중 엔진 지지를 구분합니다. 한 엔진만 유망해도 합의 부족을 실패로 간주하지 않습니다. 엔진 간 자세 차이와 생성·재예측 불일치는 별도 진단이며 자동 비결합·탈락 사유가 아닙니다.'));
  host.append(el('p','같은 후보·표적 서열만 비교하며 VHH를 따로 맞추지 않습니다. RF2/BoltzGen 재접힘 구조를 사용하는 경우 생성 단계 최종 구조와의 비교입니다. 탐색 임계값은 실험 결합에 대해 보정되지 않았습니다. 미평가를 성공으로 간주하지 않으며, 1위도 보류일 수 있습니다.'));
  host.append(el('p',(h?.label||'hotspot 미평가')+(h?.contact_fraction!=null?` · ${h.contacted}/${h.requested}개 접촉 · 기준 ${fmt(h.threshold)}`:'')));
  if(h?.status==='no_contact')host.append(el('p','신뢰도 점수가 높아도 지정 hotspot에는 접촉하지 않았습니다. 지정 부위 재현 성공으로 해석하지 마세요.'));
  host.append(el('p',f?.available?`참조 접촉 잔기 회복률 ${(100*f.reference_contact_recovery).toFixed(1)}% (${f.shared_contacts}/${f.reference_contacts}) · Jaccard ${fmt(f.contact_jaccard)} · 중원자 5 Å 기준`:'참조 결합 부위 비교 불가 · '+(f?.reason||'참조 구조 없음')));
  if(Array.isArray(f?.comparable_target_positions))host.append(el('p',`양쪽 좌표를 비교한 표적 잔기 ${f.comparable_target_positions.length}개`));
  if(f?.predicted_contacts_outside_comparable>0)host.append(el('p',`비교 가능한 좌표 범위 밖의 예측 접촉 잔기 ${f.predicted_contacts_outside_comparable}개는 제외했습니다. 회복률은 비교 가능한 구간 기준입니다.`));
  if(f?.reference_contacts_outside_comparable>0)host.append(el('p',`비교 가능한 좌표 범위 밖의 참조 접촉 잔기 ${f.reference_contacts_outside_comparable}개는 제외했습니다.`));
  const caseRow=data.cases.find(c=>c.case_id===r.case_id&&(!r.source_scope||c.source_scope===r.source_scope));
  const restrained=typeof r.restraint==='boolean'?r.restraint:caseRow?.conditions?.restraint_enabled;
  host.append(el('p',restrained===false?'후속 재예측: hotspot 위치 제약 미적용. hotspot 사용/미사용은 후보 생성 조건입니다.':restrained===true?'후속 재예측: 위치 제약 사용 기록 있음. 접촉 충족 여부는 별도로 확인하세요.':'후속 재예측 위치 제약: 기록 확인 필요.'));
  const reasons=Array.isArray(r.exclusion_reasons)?r.exclusion_reasons:String(r.exclusion_reason||'').split(';').filter(Boolean);
  host.append(el('p','실행 당시 선별: '+fmt(r.final_eligible)+(reasons.length?' · '+reasons.map(code=>reasonNames[code]||code).join(' · '):'')+'. 지정 위치 진단은 기존 선별 기록을 바꾸지 않습니다.'));
  host.append(el('p','회색은 PDB 참조, 금색은 새 VHH입니다. 서로 다른 VHH의 표적 접촉 잔기 겹침이며 같은 자세·실제 결합·KD를 입증하지 않습니다.'));
  $('binding-site-record').textContent=JSON.stringify({hotspot:h||null,reference_binding_site:f||null,pose:r.pose_assessment||null,pose_policy:data.pose_assessment?.policy||null,review_rank:r.review_rank,review_rank_scope:r.review_rank_scope},null,2);
}
function choose(r){
  selected=r;const version=++serial;
  $('selected-title').textContent=`${r.pdb_id} · hotspot ${r.arm==='on'?'사용':'미사용'} · ${r.candidate_label} · ${names[r.predictor_engine]}`;
  $('selected-record').textContent=JSON.stringify(r,null,2);$('metrics').replaceChildren();
  for(const k of ['ipsae','iptm','interface_pae','interface_plddt']){const d=el('div');d.append(el('span',captions[k]),el('strong',fmt(r[k])));$('metrics').append(d);}
  bindingSite(r);render();otherPlots(r);paePlot(r,version);$('viewer-status').textContent='원본 구조를 준비하는 중…';
  $('reference-toggle').disabled=true;$('reference-status').textContent='참조 겹침을 준비하는 중…';
  queue=queue.catch(()=>{}).then(async()=>{
    if(version!==serial)return;
    try{
      const viewer=await getViewer(),plugin=viewer.plugin;await plugin.clear();currentStructure=null;currentReference=null;currentDisplay=null;
      if(version!==serial)return;
      const asset=data.coordinates[r.structure_key];
      if(!asset){$('viewer-status').textContent='저장된 구조가 없습니다. 실패·미평가 기록을 확인하세요.';$('reference-status').textContent=referenceStatus(r);return;}
      const raw=await plugin.builders.data.rawData({data:asset.content,label:r.id}),trajectory=await plugin.builders.structure.parseTrajectory(raw,asset.format==='cif'?'mmcif':'pdb'),model=await plugin.builders.structure.createModel(trajectory);
      currentStructure=await plugin.builders.structure.createStructure(model,{name:'model',params:{}});
      const {StructureElement}=molstar.lib.structure,components=[];
      for(const [chain,color,label]of [[r.target_chain,0x3b82b2,'예측 표적'],[r.binder_chain,0xc49a40,'새 VHH']]){
        const loci=StructureElement.Loci.fromSchema(currentStructure.obj.data,{items:[{auth_asym_id:chain}]});
        const component=await plugin.builders.structure.tryCreateComponent(currentStructure,{type:{name:'bundle',params:StructureElement.Bundle.fromLoci(loci)},nullIfEmpty:true,label},chain);
        if(component){await plugin.builders.structure.representation.addRepresentation(component,{type:'cartoon',color:'uniform',colorParams:{value:color}});components.push({component,color});}
      }
      currentDisplay=new window.NanoCoFoldBenchmarkDisplay(viewer,components);
      await currentDisplay.display(displayOptions());
      const fit=data.reference_overlays?.[r.id],reference=data.reference_coordinates?.[fit?.reference_key];
      let referenceError=null;
      if(fit?.available&&reference){
        try{
          const refRaw=await plugin.builders.data.rawData({data:alignedReferencePdb(reference,fit),label:`PDB ${r.pdb_id} 참조 · 표적 Cα 정렬`});
          const refTrajectory=await plugin.builders.structure.parseTrajectory(refRaw,'pdb'),refModel=await plugin.builders.structure.createModel(refTrajectory),refStructure=await plugin.builders.structure.createStructure(refModel,{name:'model',params:{}});
          const refLoci=StructureElement.Loci.fromSchema(refStructure.obj.data,{});
          const component=await plugin.builders.structure.tryCreateComponent(refStructure,{type:{name:'bundle',params:StructureElement.Bundle.fromLoci(refLoci)},nullIfEmpty:true,label:'PDB 참조 복합체'},'pdb-reference');
          if(!component)throw Error('표시할 참조 원자 없음');
          currentReference=await plugin.builders.structure.representation.addRepresentation(component,{type:'cartoon',color:'uniform',colorParams:{value:0x697582},typeParams:{alpha:.45}});
          plugin.state.data.updateCellState(currentReference.ref,{isHidden:!$('reference-toggle').checked});
        }catch(error){referenceError=error.message;}
      }
      plugin.managers.camera.reset({durationMs:0});
      if(version===serial){
        $('viewer-status').textContent=`원본 예측 구조 · 표적 ${r.target_chain} · VHH ${r.binder_chain} · SHA256 ${r.structure_sha256}`;
        $('reference-status').textContent=referenceError?'참조 표시 오류 · '+referenceError:referenceStatus(r);
        $('reference-toggle').disabled=!currentReference;
      }
    }catch(error){if(version===serial){$('viewer-status').textContent='구조 표시 오류: '+error.message;$('reference-status').textContent='참조 표시 미완료';}}
  });
}
for(const id of ['sequence-toggle','residues-toggle'])$(id).addEventListener('change',()=>{queue=queue.catch(()=>{}).then(async()=>{if(currentDisplay)await currentDisplay.display(displayOptions());}).catch(error=>{$('viewer-status').textContent='표시 옵션 오류: '+error.message;});});
$('clear-selection').addEventListener('click',()=>{queue=queue.catch(()=>{}).then(()=>currentDisplay?.clearSelection()).catch(error=>{$('viewer-status').textContent='선택 해제 오류: '+error.message;});});
$('reference-toggle').addEventListener('change',async()=>{await queue;if(currentReference&&selected){(await viewerPromise).plugin.state.data.updateCellState(currentReference.ref,{isHidden:!$('reference-toggle').checked});$('reference-status').textContent=referenceStatus(selected);}});
$('fit').addEventListener('click',async()=>{await queue;if(viewerPromise)(await viewerPromise).plugin.managers.camera.reset({durationMs:0});});for(const [id,key]of [['target-focus','target_chain'],['binder-focus','binder_chain']])$(id).addEventListener('click',async()=>{await queue;if(!currentStructure||!selected)return;const loci=molstar.lib.structure.StructureElement.Loci.fromSchema(currentStructure.obj.data,{items:[{auth_asym_id:selected[key]}]});(await viewerPromise).plugin.managers.camera.focusLoci(loci);});$('download-structure').addEventListener('click',()=>{if(!selected)return;const asset=data.coordinates[selected.structure_key];if(!asset)return;const url=URL.createObjectURL(new Blob([asset.content],{type:'text/plain;charset=utf-8'})),a=el('a');a.href=url;a.download=selected.id+'.'+asset.format;a.click();setTimeout(()=>URL.revokeObjectURL(url),10000);});
$('source-record').textContent=JSON.stringify({summary:s,interpretation:data.interpretation,campaign_protocol:data.campaign_protocol,augmentation_protocol:data.augmentation_protocol,source_campaigns:data.source_campaigns,timing_scopes:data.timing_scopes,reference_structures:data.reference_structures,reference_sources:data.reference_sources,conditions:data.cases.map(c=>({case_id:c.case_id,source_scope:c.source_scope,conditions:c.conditions,generation_engine_info:c.generation_engine_info})),source_artifacts:data.source_artifacts},null,2);render();const first=data.models.find(r=>r.structure_key)||data.models[0];if(first)choose(first);
})();</script></body></html>'''
