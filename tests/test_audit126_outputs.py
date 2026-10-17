"""Audit 126: real score/rank/read-only export boundaries, CPU fixtures only."""
import csv
import datetime
import io
import json
import re
import zipfile
from pathlib import Path

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from nanocofold import design_benchmark_report as benchmark
from nanocofold import engine_comparison, mock, rank_select, report, score_complex
from nanocofold.config import load_config
from nanocofold.platform.results import ResultReader


def cfg():
    result = load_config()
    result["rank"].update(weights={"iptm": 1., "hotspot_contact_frac": .5}, final_n=10)
    result["rank"]["cdr3"]["per_cluster"] = 10
    return result


def population():
    return pd.DataFrame([
        {"design_id": "aa_worse", "sequence": "ACDEFG", "iptm": .8,
         "hotspot_contact_frac": .25, "hotspot_mapping_failed": False,
         "n_contacts": 5, "final_eligible": True},
        {"design_id": "zz_better", "sequence": "ACDEFG", "iptm": .800004,
         "hotspot_contact_frac": .25, "hotspot_mapping_failed": False,
         "n_contacts": 5, "final_eligible": True},
        {"design_id": "baseline", "sequence": "ACDEFG", "iptm": .2,
         "hotspot_contact_frac": .25, "hotspot_mapping_failed": False,
         "n_contacts": 5, "final_eligible": True},
    ])


def saved_run(tmp_path, *, complex_mode=False):
    tmp_path.mkdir(exist_ok=True)
    (tmp_path / "03_score").mkdir()
    config = {"run": {"mode": "complex" if complex_mode else "screen", "mock": True}}
    (tmp_path / "run_config.yaml").write_text(yaml.safe_dump(config))
    (tmp_path / "run_state.json").write_text(json.dumps({"execution": {"state": "complete", "started": 12345.5}, "phases": {"final": {}}}))
    pd.DataFrame([{"design_id": "나노바디2", "sequence": "ACDEFG", "error": "한글 검사", "=unsafe_header": "=1+1", "iptm": .30000000000000004}]).to_csv(tmp_path / "03_score/scored_candidates.csv", index=False)
    return ResultReader(tmp_path)


def test_partial_site_mapping_uses_requested_denominator_and_holds_real_rank(tmp_path):
    options = cfg()
    mock.make_mock_chai_output(tmp_path / "prediction", "test", "A" * 10, "A" * 10, [3], quality=.9)
    scored = score_complex.score_complex(tmp_path / "prediction", "A", "B", [3, 999], options)
    assert scored["hotspot_mapping_failed"] is True
    assert scored["n_hotspots"] == 1 and scored["n_hotspots_requested"] == 2
    assert scored["hotspot_contact_frac"] == scored["n_hotspots_contacted"] / 2
    frame = population()
    frame.loc[0, ["hotspot_contact_frac", "hotspot_mapping_failed"]] = [scored["hotspot_contact_frac"], True]
    ranked = rank_select.rank_candidates(frame, options)
    held = ranked.set_index("design_id").loc["aa_worse"]
    assert pd.isna(held["rank"]) and pd.isna(held["composite_score"])
    assert held["ranking_unavailable_reason"] == "weighted_metrics_unavailable"
    assert "aa_worse" not in rank_select.select_final(ranked, options, pre_ranked=True).design_id.tolist()


@pytest.mark.parametrize("failed", [True, "True", " true ", "1", 1])
def test_reload_mapping_failure_cannot_acquire_weighted_score_or_fallback(failed):
    options, frame = cfg(), population()
    frame["hotspot_mapping_failed"] = failed
    options["rank"]["fallback_to_ineligible"] = True
    ranked = rank_select.rank_candidates(frame, options)
    assert ranked["rank"].isna().all()
    assert rank_select.select_final(ranked, options, pre_ranked=True).empty
    assert any("구조 대응이 불완전" in warning for warning in rank_select.ranking_warnings(ranked, options))


def test_unweighted_mapping_issue_does_not_change_quality_eligibility():
    options, frame = cfg(), population()
    options["rank"]["weights"] = {"iptm": 1.}
    frame["hotspot_mapping_failed"] = True
    ranked = rank_select.rank_candidates(frame, options)
    assert ranked["rank"].notna().all()
    assert ranked["final_eligible"].all()


def test_failed_mapping_does_not_make_optional_cohort_metric_appear_measured():
    options, frame = cfg(), population()
    frame["hotspot_contact_frac"] = float("nan")
    frame.loc[0, ["hotspot_contact_frac", "hotspot_mapping_failed"]] = [.5, True]
    with pytest.warns(RuntimeWarning, match="hotspot_contact_frac"):
        ranked = rank_select.rank_candidates(frame, options).set_index("design_id")
    assert pd.isna(ranked.loc["aa_worse", "rank"])
    assert ranked.loc[["zz_better", "baseline"], "rank"].notna().all()
    assert "hotspot_contact_frac" in rank_select.ranking_policy(options, frame)["ignored_weights"]


def test_unrounded_composite_controls_rank_not_lexical_id_and_receipts_explain_ties():
    options, frame = cfg(), population()
    options["rank"]["weights"] = {"iptm": 1.}
    ranked = rank_select.rank_candidates(frame, options)
    assert ranked.design_id.tolist() == ["zz_better", "aa_worse", "baseline"]
    scores = ranked.set_index("design_id").composite_score
    assert scores["zz_better"] > scores["aa_worse"]
    assert round(scores["zz_better"], 4) == round(scores["aa_worse"], 4)
    frame["iptm"] = .8
    tied = rank_select.rank_candidates(frame.iloc[::-1], options)
    assert tied.design_id.tolist() == sorted(frame.design_id)
    policy = rank_select.ranking_policy(options, ranked)
    assert "unrounded_float64" in policy["score_precision"]
    assert "design_id" in policy["rank_tie_break"]
    assert "do not establish a scientific difference" in policy["rank_interpretation"]


def test_comparison_plate_ids_disambiguate_engine_without_cross_engine_order(tmp_path):
    options, frame = cfg(), population()
    options["screen"]["comparison"]["enabled"] = True
    rows = pd.concat([frame.assign(predictor_engine=engine, source_sequence_id=frame.design_id)
                      for engine in ("boltz2", "chai1")], ignore_index=True)
    ranked, selected = engine_comparison.rank_and_select(rows, options)
    assert not selected.duplicated(["plate_id", "plate_position"]).any()
    assert set(selected.engine_plate_id) == {"plate_01"}
    assert set(selected.plate_id) == {"boltz2__plate_01", "chai1__plate_01"}
    for _, group in ranked.groupby("predictor_engine"):
        assert sorted(group["rank"].tolist()) == [1, 2, 3]
    receipt = engine_comparison.write_synthesis_by_engine(selected, options, tmp_path)
    assert receipt["requires_researcher_selection"] is True
    assert pd.read_csv(tmp_path / "synthesis_plate_list.csv").empty


def test_pae_palette_uses_actual_vendor_direction_consistently():
    root = Path(benchmark.__file__).parent
    bundled = (root.parent / "assets/plotly.min.js").read_text()
    match = re.search(r'Blues:\[\[0,"rgb\(([^)]*)\)"\].*?\[1,"rgb\(([^)]*)\)"\]\]', bundled)
    assert match, "Verify the actual bundled Plotly palette rather than assuming standard Blues"
    low, high = [tuple(map(int, channel.split(","))) for channel in match.groups()]
    assert sum(low) < sum(high), "This bundle stores the dark endpoint at zero"
    explorer = (root / "templates/platform.js").read_text()
    assert re.search(r'colorscale:\s*"Blues",\s*zmin:\s*0,\s*zmax:\s*30,\s*reversescale:\s*true', explorer)
    rendered = (root / "design_benchmark_report.py").read_text()
    assert "zmin:0,zmax:30,colorscale:'Blues',reversescale:true" in rendered


@pytest.mark.parametrize("value,expected", [(.30000000000000004, "0.3"), (-0., "0"), (72.96000000000001, "72.96"), (1e-8, "1e-08")])
def test_readable_display_never_rounds_small_nonzero_to_zero(value, expected):
    assert report.display_number(value) == expected


def test_real_report_display_utc_units_and_source_preservation(tmp_path):
    options = cfg()
    options["report"]["interactive_structure"] = False
    options["target"]["hotspots"] = []
    rows = pd.DataFrame([{"design_id": "original", "sequence": "ACDEFG", "iptm": .30000000000000004,
                          "composite_score": -0., "interface_plddt": 72.96000000000001,
                          "interface_pae": 1e-8, "final_eligible": True}])
    before = rows.copy(deep=True)
    path = report.generate_report(tmp_path / "absent-source", options, {}, scored=rows, out_path=tmp_path / "report.html")
    text = path.read_text()
    assert "72.96000000000001" not in text and "0.30000000000000004" not in text
    assert re.search(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", text)
    assert "계면 PAE (Å)" in text
    assert dict(report.HOTSPOT_COLUMNS)["dist_to_glycan"] == "글리칸 거리 (Å)"
    pd.testing.assert_frame_equal(rows, before)


def test_all_csv_headers_cells_and_excel_korean_round_trip(tmp_path):
    reader = saved_run(tmp_path)
    source = (tmp_path / "03_score/scored_candidates.csv").read_bytes()
    csv_data = reader.export("csv")[0]
    assert csv_data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(csv_data.decode("utf-8-sig"))))
    assert "'=unsafe_header" in rows[0] and rows[0]["'=unsafe_header"] == "'=1+1"
    assert rows[0]["candidate_id"] == "나노바디2"
    benchmark_data = benchmark._text_data([{"=formula": "한글 검사", "note": "=1+1"}])
    benchmark_rows = list(csv.DictReader(io.StringIO(benchmark_data.decode("utf-8-sig"))))
    assert benchmark_rows == [{"'=formula": "한글 검사", "note": "'=1+1"}]
    delivery = reader.export("report_zip")[0]
    with zipfile.ZipFile(io.BytesIO(delivery)) as archive:
        header = next(csv.reader(io.StringIO(archive.read("candidates.csv").decode("utf-8-sig"))))
        assert "'=unsafe_header" in header and "=unsafe_header" not in header
    assert (tmp_path / "03_score/scored_candidates.csv").read_bytes() == source
    # The canonical machine-interchange CSV is deliberately not changed to BOM.
    assert not source.startswith(b"\xef\xbb\xbf")


@pytest.mark.parametrize("kind", ["csv", "report_zip", "benchmark", "atomic"])
def test_colliding_formula_safe_headers_fail_without_losing_original_columns(tmp_path, kind):
    from nanocofold.utils.atomic import write_csv
    frame = pd.DataFrame([{"design_id": "original", "sequence": "ACDEFG", "=extra": "first", "'=extra": "second"}])
    before = frame.copy(deep=True)
    if kind in {"csv", "report_zip"}:
        reader = saved_run(tmp_path)
        path = tmp_path / "03_score/scored_candidates.csv"
        frame.to_csv(path, index=False)
        source = path.read_bytes()
        with pytest.raises(ValueError, match="열 이름이 중복"):
            reader.export(kind)
        assert path.read_bytes() == source
    elif kind == "benchmark":
        with pytest.raises(ValueError, match="열 이름이 중복"):
            benchmark._text_data(frame.to_dict(orient="records"))
    else:
        path = tmp_path / "existing.csv"
        path.write_bytes(b"previous valid output\n")
        with pytest.raises(ValueError, match="열 이름이 중복"):
            write_csv(path, frame)
        assert path.read_bytes() == b"previous valid output\n"
    pd.testing.assert_frame_equal(frame, before)


def test_xlsx_exports_are_byte_reproducible_and_keep_run_times_text_and_science(tmp_path, monkeypatch):
    import openpyxl.writer.excel as writer
    from types import SimpleNamespace
    reader = saved_run(tmp_path)
    class Early(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2020, 1, 2, tzinfo=tz)
    class Late(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2030, 9, 10, tzinfo=tz)
    monkeypatch.setattr(writer, "datetime", SimpleNamespace(datetime=Early, timezone=datetime.timezone))
    early = reader.export("xlsx")[0]
    monkeypatch.setattr(writer, "datetime", SimpleNamespace(datetime=Late, timezone=datetime.timezone))
    late = reader.export("xlsx")[0]
    assert early == late
    book = load_workbook(io.BytesIO(late))
    values = list(book["All candidates"].values)
    assert "=unsafe_header" in values[0]
    assert "나노바디2" in values[1]
    assert .30000000000000004 == pytest.approx(next(value for value in values[1] if isinstance(value, float)))
    assert "12345.5" in " ".join(str(cell.value) for row in book["Provenance"] for cell in row)
    with zipfile.ZipFile(io.BytesIO(late)) as archive:
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())
        assert b"2030-09-10" not in archive.read("docProps/core.xml")
        assert b"1980-01-01T00:00:00Z" in archive.read("docProps/core.xml")
        assert not any(b"<f>" in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/"))


def test_benchmark_xlsx_headers_are_text_and_metadata_is_deterministic():
    payload = {key: [] for key in ("models", "generated_candidates", "cases", "targets", "timing_summary", "source_artifacts")}
    payload["models"] = [{"=unsafe_header": "한글 검사", "score": .123456789012345}]
    first = benchmark._workbook(payload)
    second = benchmark._workbook(payload)
    assert first == second
    book = load_workbook(io.BytesIO(first))
    assert book["Models"]["A1"].value == "=unsafe_header" and book["Models"]["A1"].data_type == "s"
    assert book["Models"]["A2"].value == "한글 검사"
    assert book["Models"]["B2"].value == .123456789012345


def test_imported_complex_mock_summary_keeps_recorded_mock_warning(tmp_path):
    reader = saved_run(tmp_path, complex_mode=True)
    assert reader.summary()["mock"] is True
